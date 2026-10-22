"""Render the completed Chinese Markdown report as a plain, editable Word report."""
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs/CLE_Dgate全量四臂三张上限实验结果_20261019.md'
TARGET = SOURCE.with_suffix('.docx')


def font(run, size=10.5, bold=False):
    run.font.name = 'Microsoft YaHei'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


def text(paragraph, value, size=10.5, bold=False):
    # Break long artifact paths without changing their visible content.
    value = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', value)
    value = value.replace('`', '').replace('**', '')
    value = re.sub(r'\S{45,}', lambda m: re.sub(r'([/_])', lambda s: s[0] + '\u200b', m[0]), value)
    font(paragraph.add_run(value), size, bold)


def main():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin, section.right_margin = Cm(1.8), Cm(1.8)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Microsoft YaHei', Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05
    for style in document.styles:
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    for name, size in [('Title', 19), ('Heading 1', 14), ('Heading 2', 11.5)]:
        style = document.styles[name]
        style.font.name, style.font.size = 'Microsoft YaHei', Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text(footer, 'CLE Dgate 四臂实验报告  ·  ', 8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', c) for c in cells):
                    rows.append(cells)
                i += 1
            columns = len(rows[0])
            table = document.add_table(rows=len(rows), cols=columns)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            widths = [17.4 / columns] * columns
            if columns == 7:
                widths = [2.6] + [14.8 / 6] * 6
            if columns == 6:
                widths = [4.2] + [13.2 / 5] * 5
            for col, width in zip(table.columns, widths):
                col.width = Cm(width)
            borders = OxmlElement('w:tblBorders')
            for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                edge = OxmlElement('w:' + side)
                for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    edge.set(qn('w:' + key), val)
                borders.append(edge)
            table._tbl.tblPr.append(borders)
            for r, values in enumerate(rows):
                row = table.rows[r]
                no_split = OxmlElement('w:cantSplit')
                row._tr.get_or_add_trPr().append(no_split)
                if r == 0:
                    row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for c, value in enumerate(values):
                    cell = row.cells[c]
                    cell.width = Cm(widths[c])
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    properties = cell._tc.get_or_add_tcPr()
                    margins = OxmlElement('w:tcMar')
                    for side in ['top', 'bottom', 'left', 'right']:
                        margin = OxmlElement('w:' + side)
                        margin.set(qn('w:w'), '70')
                        margin.set(qn('w:type'), 'dxa')
                        margins.append(margin)
                    properties.append(margins)
                    if r == 0:
                        fill = OxmlElement('w:shd')
                        fill.set(qn('w:fill'), 'F2F2F2')
                        properties.append(fill)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.line_spacing = 1.05
                    p.paragraph_format.keep_with_next = r < len(rows) - 1
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    text(p, value, 8.5 if columns >= 7 else 9, r == 0)
            spacer = document.add_paragraph()
            spacer.paragraph_format.space_after = Pt(0)
            spacer.paragraph_format.line_spacing = Pt(3)
            font(spacer.add_run(''), 3)
            continue
        if line.startswith('# '):
            p = document.add_paragraph(style='Title')
            title = line[2:].replace('D_gate', 'Dgate')
            text(p, title, 19, True)
        elif line.startswith('## '):
            p = document.add_paragraph(style='Heading 1')
            text(p, line[3:], 14, True)
        elif line.startswith('### '):
            p = document.add_paragraph(style='Heading 2')
            text(p, line[4:].replace('：', '  '), 11.5, True)
        else:
            p = document.add_paragraph()
            if line.startswith('- '):
                p.paragraph_format.left_indent = Cm(.35)
                p.paragraph_format.first_line_indent = Cm(-.25)
                line = '• ' + line[2:]
            text(p, line)
        i += 1
    document.core_properties.title = 'CLE Dgate全量四臂三张上限实验结果'
    document.core_properties.subject = '298题旧版CLE P7冻结Q四臂实验与逐题诊断'
    document.save(TARGET)
    print(TARGET)


if __name__ == '__main__':
    main()
