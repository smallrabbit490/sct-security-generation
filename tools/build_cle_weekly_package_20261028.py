"""用首周已保存材料生成三种阅读视图；不连接工作台、不调用实验入口。"""
from pathlib import Path
import json
import re
import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET
from urllib.parse import urlparse, unquote

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from tools import build_cle_research_stage_map as diagram
from tools.build_cle_horizontal_ideas import wrap_clean

OUT = ROOT / 'docs/research/reports/CLE-PKG-20261014-20261021-v01'
OLD = ROOT / 'docs/research/reports/CLE-RPT-20261014-20261021-weekly-v06/CLE-RPT-20261014-20261021-weekly-v06.md'
SUMMARY = '10月14日-10月21日本周结论与下一步'
MAP = '10月14日-10月21日思路脉络图'
REPORT = '10月14日-10月21日实验探索报告'
NODE = Path('C:/Users/fujunye/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe')
SHARP = NODE.parents[1] / 'node_modules/sharp'


def write(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.strip() + '\n', encoding='utf-8')


def mixed(draw, xy, value, size=27, fill='#263648', bold=False):
    """汉字宋体，拉丁字符及数字Times New Roman；按字符宽度连续排版。"""
    x, y = xy
    for part in re.findall(r'[\x00-\x7f]+|[^\x00-\x7f]+', str(value)):
        latin = all(ord(c) < 128 for c in part)
        font = ImageFont.truetype('C:/Windows/Fonts/' + ('timesbd.ttf' if bold else 'times.ttf') if latin else 'C:/Windows/Fonts/simsun.ttc', size)
        draw.text((x, y), part, font=font, fill=fill)
        x += draw.textlength(part, font=font)
    return x


def bars(path, title, labels, values, suffix='', maxval=None, foot='', height=None):
    width = 1440
    height = height or (115 + len(labels) * 65 + 65)
    im = Image.new('RGB', (width, height), 'white'); dr = ImageDraw.Draw(im)
    mixed(dr, (30, 15), title, 31, bold=True)
    top = 80; maximum = maxval or max(values) * 1.14
    for i, (label, value) in enumerate(zip(labels, values)):
        y = top + i * 65
        mixed(dr, (32, y + 4), label, 26)
        x1 = 410; x2 = x1 + 760 * value / maximum
        dr.rounded_rectangle((x1, y, x2, y + 36), radius=3, fill='#315a77')
        mixed(dr, (x2 + 16, y + 3), f'{value}{suffix}', 27)
    if foot: mixed(dr, (32, height - 42), foot, 23, '#526071')
    im.save(path, dpi=(220, 220))


def figures():
    fig = OUT / 'figures'; fig.mkdir(exist_ok=True)
    bars(fig/'a1_coverage.png', '各历史快照获得有效反馈的经验卡',
         ['增长集观察学习后', '资格清理后', '补充32题后', '初始化集限三卡后', '初始化集不限三卡后', '在线初始化集后', '期末增长集23题后'],
         [466,416,436,669,926,967,972], maxval=1100,
         foot='卡库共1939张；不同配置快照分别展示，非同一方案逐轮收益。')
    bars(fig/'a2_gate.png', '共同完整测量的296题中，功能与安全同时通过的题数',
         ['原CLE选卡', '相似度选卡', '学习Q选卡', '打乱Q选卡'], [96,103,102,100], maxval=125,
         foot='各组均为同一296题；两道缺测不参与这张配对图。')
    bars(fig/'a3_selection.png', '选卡变化与采用新卡子集的结果',
         ['可比较题目', '选入卡编号改变', '采用新增卡', '该子集原先Joint', '该子集后来Joint', '该子集救回', '该子集损失'],
         [296,281,225,103,110,16,9], maxval=340,
         foot='Joint表示功能与安全同时通过；采用新卡子集是事后分组。')
    bars(fig/'a4_losses.png', '18道损失题的后续代码复盘',
         ['重定义已有配置', '公开契约缺失', '资源策略', '网络问题'], [13,2,2,1], maxval=16,
         foot='后续复盘仅解释旧失败记录，局部修复不改原模型成绩。')
    bars(fig/'a5_labels.png', '914个题卡配对的AI初标',
         ['直接适用', '有条件适用', '无关', '冲突', '信息不足'], [89,385,285,153,2], maxval=450,
         foot='人工确认0对；图中分布不是正式检索准确率。')
    # 摘要图保留唯一主要比较，读者无需先理解四个实验代号。
    bars(fig/'summary_key.png', '功能与安全同时通过的题数',
         ['按相似度选卡', '加入经验评分'], [103,102], maxval=125,
         foot='同一296题；通过指代码同时满足功能与安全要求。')


IDEAS = [
    dict(id='r1', title='Q会改变选卡，也能提高正确率吗？',
         answer='Q改变了实际选入的卡；本次旧冻结Q没有稳定提高正确率。', judgment='选卡作用得到支持；稳定提分未获支持',
         steps=[('研究问题','评分能否帮助模型选到更有用的经验，并写出更好的代码？','r1'),
                ('验证路径','先重复选卡，区分换卡作用与模型波动；再固定评分比较代码。','a3'),
                ('主要证据','281/296题换卡。共同296题中，学习Q通过102题，相似度通过103题。','a2'),
                ('后续判断','机制已作用到选卡；下一步用强化后的冻结Q检验跨题族效果。','t2')],
         links=[('A3 选卡变化','a3'),('A2 冻结对照','a2'),('T2 强化后评价','t2')]),
    dict(id='r2', title='让更多卡获得反馈，做题是否改善？',
         answer='更多卡获得了反馈，但覆盖增加没有自动转化为更好的同题成绩。', judgment='覆盖扩大得到支持；覆盖等于收益不成立',
         steps=[('研究问题','反馈不足是否限制评分作用？扩大来源题和使用机会后会怎样？','r2'),
                ('验证路径','对比不同历史学习配置，记录有反馈的卡数，再看冻结同题回测。','a1'),
                ('主要证据','有反馈卡466→972张。同题回测109→100题通过；两阶段选择方式不同。','a1'),
                ('期末安排','在线增长集完成23/273题；保留剩余250题，再冻结并评价。','t1')],
         links=[('A1 覆盖与回测','a1'),('A6 期末状态','a6'),('T1 剩余学习','t1')]),
    dict(id='r3', title='代码为什么从通过变成失败？',
         answer='后续复盘定位了具体失败路径：18题中13题重定义已有配置。', judgment='具体代码原因已定位；Q因果贡献未单独分离', later=True,
         steps=[('研究问题','选卡变化后出现损失，问题来自经验内容、题目条件还是实现？','r3'),
                ('验证路径','对齐18道损失题的题面、卡、前后代码与验证；做局部零API诊断。','a4'),
                ('主要证据','13题配置、2题契约、2题资源、1题网络；前两类局部诊断恢复。','a4'),
                ('解释与影响','重点检查来源假设和公开要求。局部修复不替代原模型成绩。','a4')],
         links=[('A4 代码与契约复盘','a4'),('A2 原始代码成绩','a2')]),
    dict(id='r4', title='检索到的经验是否真正适用？',
         answer='已完成914个题卡配对的AI初标；正式检索四指标待人工确认。', judgment='初标材料已完成；检索质量尚待检验',
         steps=[('研究问题','经验能否用于当前题？选入是否准确，又遗漏了多少适用经验？','r4'),
                ('验证路径','构建共同候选池，按五类初标，隐藏方法和排名后供人工复核。','a5'),
                ('主要证据','89直接适用、385有条件适用、285无关、153冲突、2信息不足。','a5'),
                ('期末安排','完成T3人工裁决后计算候选池内四指标；全库召回仍需全库标签。','t3')],
         links=[('A5 标签与口径','a5'),('T3 人工复核','t3')]),
]


def maps():
    folder=OUT/'diagrams';folder.mkdir(exist_ok=True); diagram.OUT=folder
    for j,spec in enumerate(IDEAS,1):
        d=diagram.Diagram(1500,650,spec['title'])
        d.background(0,0,1500,650,'#ffffff','none',0)
        d.text(30,47,spec['title'],34,bold=True)
        d.rect(30,76,1440,92,'#edf3f7','#b8c9d6',7)
        d.text(50,113,'当前判断：'+spec['answer'],28,bold=True)
        d.text(50,147,'证据判断：'+spec['judgment'],22,color='#415467')
        for i,(heading,body,anchor) in enumerate(spec['steps']):
            x=30+i*370
            href='../'+REPORT+'.md#'+anchor
            d.rect(x,216,330,265,'#f8fafc','#becbd6',7,id=f'box{i}',href=href)
            d.text(x+18,253,heading,27,bold=True)
            d.text(x+18,299,'\n'.join(wrap_clean(body,294,26)),26)
            if i<3:d.line([(x+330,349),(x+365,349)],'#6b7f91',source=f'box{i}',target=f'box{i+1}')
        note='后续复盘：分析发生在10月21日之后，保留原结果，不回填首周期末判断。' if spec.get('later') else '图示范围：首周历史材料；期末安排截至10月21日22:26。'
        d.text(30,531,note,21,color='#526071')
        states=['执行：选卡与旧Q对照已完成；强化后评价待做。',
                '执行：历史学习与同题回测已完成；在线增长学习期末暂停。',
                '执行：18题后续复盘已完成；分析时间与首周实验时间分开。',
                '执行：AI初标已完成；人工确认与正式四指标待做。']
        d.text(30,570,states[j-1],21,color='#526071')
        x=30
        for label,anchor in spec['links']:
            d.text(x,618,'【'+label+'】',23,color='#255c87',href='../'+REPORT+'.md#'+anchor);x+=440
        d.export(f'思路{j}')
        svg=folder/f'思路{j}.svg'
        ET.register_namespace('', 'http://www.w3.org/2000/svg')
        tree=ET.parse(svg); tree.getroot().find('{http://www.w3.org/2000/svg}desc').text='首周横向研究问题、验证、证据与下一步；支持状态与执行状态分开。'
        tree.write(svg,encoding='utf-8',xml_declaration=True)
    code="const fs=require('fs'),sharp=require(process.argv[1]);(async()=>{for(const f of JSON.parse(process.argv[2]))await sharp(fs.readFileSync(f),{density:144}).png().toFile(f.replace(/\\.svg$/,'.png'));})().catch(e=>{process.stderr.write(String(e));process.exit(1)});"
    subprocess.run([str(NODE),'-e',code,str(SHARP),json.dumps([str(folder/f'思路{i}.svg') for i in range(1,5)])],check=True)
    for i in range(1,5):
        # 文档已有可跳转标题和链接行；嵌图只保留判断与路径，避免重复印两遍。
        with Image.open(folder/f'思路{i}.png') as im:
            im.crop((0,130,3000,1176)).save(folder/f'思路{i}_报告.png',dpi=(144,144))
    chunks=['# '+MAP,'[一页结论]('+SUMMARY+'.md)　[完整周报]('+REPORT+'.md)',
            '四张横版图各对应一条问题。点击下方编号可直接进入实验；SVG保留图内链接。图中箭头表示论证路径，执行顺序以当周实验安排为准。']
    for i,spec in enumerate(IDEAS,1):
        if i>1:chunks.append('<!-- pagebreak -->')
        chunks += [f'<a id="{spec["id"]}"></a>',f'## {spec["title"]}',
                   f'![{spec["answer"]}](diagrams/思路{i}_报告.png)',
                   '　'.join(f'[【{label}】]({REPORT}.md#{anchor})' for label,anchor in spec['links'])+
                   f'　[打开可点击横版图](diagrams/思路{i}.svg)',
                   '[返回一页结论]('+SUMMARY+'.md)']
    write(OUT/(MAP+'.md'),'\n\n'.join(chunks))


def report():
    source=OLD.read_text(encoding='utf-8-sig')
    value=source
    value=value.replace('# '+REPORT,'# '+REPORT+'\n\n[先看一页结论]('+SUMMARY+'.md)　[查看横版思路图]('+MAP+'.md)',1)
    # 保留原数字、表格、步骤与历史待做；各实验先给直接答案。
    answers={
      'a1':'**先看结论：** 反馈覆盖从24.03%扩大到50.13%，同题回测却从109题通过降到100题。扩大覆盖已经实现，提分没有随之实现。',
      'a2':'**先看结论：** 在共同完整测量的296题上，学习Q比相似度少通过1题，比原CLE多6题，比打乱Q多2题。配对差区间均跨过0，本次对照未建立稳定提分。',
      'a3':'**先看结论：** 296道可比较题中，281题改变了选入卡编号，225题采用了新增卡。Q的影响到达了实际选卡与生成过程。',
      'a4':'**先看结论：** 后续复盘发现，18道损失题中13题重定义已有配置，另5题分属契约、资源和网络问题。本节解释历史失败，不改写首周正式成绩。',
      'a5':'**先看结论：** 914个题卡配对已有AI初标，人工确认仍为0；完成标签裁决是形成正式检索四指标的下一步。',
      'a6':'**先看结论：** 期末在线链完成初始化集303题及增长集23题；增长集余250题，强化后的评价尚未执行。',
    }
    titles={
      'a1':'【A1】反馈覆盖扩大后，同题成绩是否提高？',
      'a2':'【A2】旧冻结Q能否提高全评价集的通过率？',
      'a3':'【A3】Q是否改变实际卡片编号并进入生成过程？',
      'a4':'【A4】18道损失题为什么失败？（后续复盘）',
      'a5':'【A5】已有标签能否给出正式检索质量？',
      'a6':'【A6】期末在线学习实际完成到了哪里？',
    }
    for anchor,answer in answers.items():
        pattern=r'(<a id="'+anchor+r'"></a>\s*)### [^\n]+'
        value=re.sub(pattern,lambda m:m[1]+'### '+titles[anchor]+'\n\n'+answer,value)
    value=value.replace('## 二、截至本周期末的待做实验','<!-- pagebreak -->\n## 二、截至本周期末的待做实验')
    value=value.replace('![图2 旧冻结Q全D_gate结果；Q2有2题未完整测量，未计为通过]',
                        '![图2 在共同完整测量的296题上比较Joint通过题数]')
    value=value.replace('| Q2 学习Q | 296 | 118 | 228 | 102（34.23%） |',
                        '| Q2 学习Q | 296 | 118 | 228 | 102（已测34.46%） |')
    value=value.replace('共同296题上，Q2相对Q0',
      '表中Q0、Q1、Q3通过率的分母为298；Q2已测通过率为102/296。原记录的102/298＝34.23%是已确认通过数占计划量，另2题为缺测，不计作已测失败。图2统一采用共同296题：Q0为96题，Q1为103题，Q2为102题，Q3为100题。\n\n共同296题上，Q2相对Q0')
    value=value.replace('图2中学习Q比原CLE和打乱Q略高，比纯相似度略低，且同题救回与损失相互抵消。因此，本次旧对照没有建立学习Q的稳定增益。它覆盖298题名单，但学习Q组有2题缺测，是带缺测的全量名单证据，不能被小样本较好结果替换；也不能证明所有Q学习方法都无效。',
      '图2显示，在同一批可比较题上，学习Q的收益与损失大幅抵消，未形成稳定优势。因此，本次结果支持Q能够改变选择，但不支持采用这份旧冻结Q来稳定提分。')
    value=value.replace('Q0的候选边界与其他组可能不同；与Q0比较时，不能把所有差异都归因于Q排序。',
                        'Q0保留原候选路径；对Q排序的解释主要依靠与相似度及打乱Q的比较，Q0作为原流程参照。')
    value=value.replace('**回答：** 对18道损失题的后续代码复盘发现，', '**回答（后续复盘）：** 对18道损失题的代码复盘发现，')
    # 从原目录移入同深度新目录后，仅修复外部报告链接；原生证据相对路径保持有效。
    write(OUT/(REPORT+'.md'),value)
    # 保留首周全部数值；原Q2百分数改为解释旧口径，新增已测口径，不修改计数。
    old_nums=set(re.findall(r'\d[\d,]*(?:\.\d+)?',source))
    new_nums=set(re.findall(r'\d[\d,]*(?:\.\d+)?',value))
    missing=sorted(old_nums-new_nums)
    if missing:raise ValueError(('historical_numbers_missing',missing))


def summary():
    text=f'''# 10月14日—21日 本周结论与下一步

本周研究：给经验打分，能否帮助模型写出既正确又安全的代码？Q就是根据历史做题反馈形成的经验评分。

**主要发现：** [Q确实改变了实际选用的经验]({REPORT}.md#a3)，但[旧评分没有稳定提高正确率]({REPORT}.md#a2)。在同一批可比较题上，评分选卡通过102题，相似度选卡通过103题。扩大反馈让更多经验获得评分，也没有自动改善同题表现。

**对主线的影响：** 后续判断应落在“哪些经验真正帮助解题”，同时检查经验是否适用、模型是否正确落实；卡片被选中或获得评分，都不能代替代码结果。

![关键对照结果](figures/summary_key.png)

**下一步：** 按本周期末安排，先补齐在线学习，再固定评分检验迁移效果；完成适用性标签复核，用于判断检索是否找对经验。

[查看横版思路图]({MAP}.md)　[查看完整周报]({REPORT}.md)
'''
    write(OUT/(SUMMARY+'.md'),text)
    # 图片文字单独计入；链接计可见标签，路径不计。
    visible=re.sub(r'!\[([^\]]*)\]\([^)]*\)',r'\1',text)
    visible=re.sub(r'\[([^\]]*)\]\([^)]*\)',r'\1',visible)
    visible=re.sub(r'[#*]','',visible)
    graph='功能与安全同时通过的题数按相似度选卡103加入经验评分102同一296题；通过指代码同时满足功能与安全要求。'
    count=sum(not c.isspace() for c in visible+graph)
    if count>500:raise ValueError(('summary_too_long',count))
    return count


def setfont(style,size=None):
    style.font.name='Times New Roman'
    if size:style.font.size=Pt(size)
    style.font.color.rgb=RGBColor(0,0,0)
    pr=style.element.get_or_add_rPr();rf=pr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');pr.insert(0,rf)
    for k in list(rf.attrib):
        if 'Theme' in k or 'theme' in k:del rf.attrib[k]
    for k,v in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','宋体'),('cs','Times New Roman')]:rf.set(qn('w:'+k),v)


def hyperlink(p,label,target,base):
    h=OxmlElement('w:hyperlink');h.set(qn('w:history'),'1')
    path,_,anchor=target.partition('#')
    if path:
        if urlparse(path).scheme in ('https','http'):url=target
        else:
            f=(base/unquote(path)).resolve()
            if f.suffix=='.md' and f.with_suffix('.docx').exists():f=f.with_suffix('.docx')
            # 三个本轮交付物先行建立目标名，构建顺序不影响Word链接类型。
            if f.parent==base and f.suffix=='.md':f=f.with_suffix('.docx')
            url=f.as_uri()
        h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    if anchor:h.set(qn('w:anchor'),anchor)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    rf=OxmlElement('w:rFonts')
    for k,v in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','宋体')]:rf.set(qn('w:'+k),v)
    pr.append(rf);u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u)
    color=OxmlElement('w:color');color.set(qn('w:val'),'255C87');pr.append(color)
    r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)


def inline(p,text,base):
    for part in re.split(r'(\[[^\]]+\]\([^)]+\)|\*\*.+?\*\*|`[^`]+`)',text):
        if not part:continue
        m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',part)
        if m:hyperlink(p,m[1],m[2],base)
        elif part.startswith('**'):p.add_run(part[2:-2]).bold=True
        else:p.add_run(part.strip('`') if part.startswith('`') else part)


def build_word(md,kind='report'):
    doc=Document();sec=doc.sections[0]
    landscape=kind=='map'
    sec.page_width=Cm(29.7 if landscape else 21);sec.page_height=Cm(21 if landscape else 29.7)
    sec.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sec.left_margin=sec.right_margin=Cm(1.7 if landscape else 2)
    sec.top_margin=sec.bottom_margin=Cm(1.5 if landscape else 1.8)
    for st in doc.styles:
        if st.type in (1,2):setfont(st)
        for b in list(st.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
    size=12 if kind=='summary' else 11
    for name,n in [('Normal',size),('Title',19),('Heading 1',15),('Heading 2',13),('Heading 3',11.5),('Caption',9.5)]:
        setfont(doc.styles[name],n)
        pf=doc.styles[name].paragraph_format;pf.first_line_indent=Pt(0);pf.space_after=Pt(6)
        if name.startswith('Heading'):pf.keep_with_next=True;pf.space_before=Pt(10)
    doc.styles['Caption'].font.italic=False;doc.styles['Caption'].font.bold=False
    pf=doc.styles['Normal'].paragraph_format;pf.line_spacing=1.16 if kind=='report' else 1.2;pf.widow_control=True
    pf.first_line_indent=Pt(22 if kind=='report' else 0)
    if kind!='summary':
        p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run('第 ');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f);p.add_run(' 页')
    text=md.read_text(encoding='utf-8');lines=text.splitlines();i=0;anchor=None;bid=0;pagebreak=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line=='<!-- pagebreak -->':pagebreak=True;continue
        m=re.fullmatch(r'<a id="([^"]+)"></a>',line)
        if m:anchor=m[1];continue
        if line.startswith('<!--'):continue
        m=re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)',line)
        if m:
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            width=Cm(26.3 if landscape else 17)
            p.add_run().add_picture(str(md.parent/m[2]),width=width)
            if kind=='report':
                c=doc.add_paragraph(m[1],'Caption');c.alignment=WD_ALIGN_PARAGRAPH.CENTER
            continue
        if line.startswith('|'):
            raw=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):raw.append(lines[i].strip());i+=1
            rows=[[v.strip() for v in r.strip('|').split('|')] for r in raw if not re.fullmatch(r'[| :\-]+',r)]
            table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Table Grid'
            for j,row in enumerate(rows):
                cells=table.add_row().cells;tr=table.rows[-1]._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
                if j==0:tr.append(OxmlElement('w:tblHeader'))
                for cell,val in zip(cells,row):
                    p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                    inline(p,val,md.parent)
                    for r in p.runs:r.font.size=Pt(9.5);r.bold=j==0
                    if j==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDF3F7');cell._tc.get_or_add_tcPr().append(shade)
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
            continue
        m=re.match(r'^(#{1,4}) (.*)',line)
        if m:
            n=len(m[1]);p=doc.add_paragraph(style='Title' if n==1 else 'Heading '+str(n-1));inline(p,m[2],md.parent)
            if n==1:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        else:
            p=doc.add_paragraph();inline(p,line,md.parent)
            if re.match(r'^\d+\.',line) or line.startswith(('**先看','[','原始证据：','实验导航：','当周状态：','实验步骤：')):
                p.paragraph_format.first_line_indent=Pt(0)
            if line.startswith(('[返回','原始证据：')):
                for r in p.runs:r.font.size=Pt(9)
            if line=='实验步骤：':p.paragraph_format.keep_with_next=True
        if pagebreak:p.paragraph_format.page_break_before=True;pagebreak=False
        if anchor:
            bid+=1;st=OxmlElement('w:bookmarkStart');st.set(qn('w:id'),str(bid));st.set(qn('w:name'),anchor)
            en=OxmlElement('w:bookmarkEnd');en.set(qn('w:id'),str(bid));p._p.insert(0,st);p._p.append(en);anchor=None
    doc.core_properties.title=lines[0].lstrip('# ');doc.core_properties.author='CLE项目研究记录'
    out=md.with_suffix('.docx');doc.save(out);return out


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    figures();report();maps();count=summary()
    paths=[build_word(OUT/(REPORT+'.md')),build_word(OUT/(MAP+'.md'),'map'),build_word(OUT/(SUMMARY+'.md'),'summary')]
    write(OUT/'build_manifest.json',json.dumps({'summary_visible_characters':count,'source_report':str(OLD.relative_to(ROOT)),
      'period_end':'2026-10-21T22:26:00+08:00','historical_numeric_tokens_retained':True,
      'report_sections':['A1','A2','A3','A4','A5','A6'],'historical_todos':['T1','T2','T3'],
      'later_review_explicit':'A4','generated_docx':[str(p) for p in paths]},ensure_ascii=False,indent=2))
    print(json.dumps({'characters':count,'docx':[str(p) for p in paths]},ensure_ascii=False))


if __name__=='__main__':
    import argparse
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--source',type=Path)
    p.add_argument('--kind',choices=['summary','map','report'],default='report')
    args=p.parse_args()
    if args.source:print(build_word(args.source.resolve(),args.kind))
    else:main()
