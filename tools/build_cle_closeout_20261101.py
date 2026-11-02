"""汇入已完成的两批原生结果；不调用模型、不改变实验状态。"""
from pathlib import Path
from collections import Counter
import json
import re

ROOT = Path(__file__).resolve().parents[1]
RUN = Path('E:/ResearchRuns/CLE/20261027_full_requirements')
Q = RUN/'probe_relay_20261028/Q.retest/model_switch_flash_20261101'
G = RUN/'FG_LABELS/model_switch_flash_20261101'
OUT = ROOT/'docs/research/reports/CLE-RPT-20261101-two-batch-closeout-v01'
MD = OUT/'11月1日Q回测与记忆标签结果分析.md'
NAMES = {'1017_only':'原1017模型','flash_only':'无1017模型','mixed_saved_prefix':'混合旧前缀'}
LEVELS = {'L0':'无卡','L1':'风险提醒','L2':'原则与适用条件','L3':'原则加目标操作模板','L4':'原则加来源示例'}


def read(path):
    return json.loads(path.read_text(encoding='utf-8'))


def table(headers, rows):
    return '\n'.join(['| '+' | '.join(map(str, headers))+' |',
        '| '+' | '.join(['---']*len(headers))+' |']+
        ['| '+' | '.join(map(str, row))+' |' for row in rows])


def pp(value):
    return '无可测配对' if value is None else f'{100*value:+.2f}'


def clarify(text, q, g):
    """落实已保存封闭读者清单；原统计和原运行记录不变。"""
    from collections import defaultdict
    roots = {'Q': Q.parent, 'G': G.parent}
    rows = {}
    for kind, root in roots.items():
        segments = read(root/'execution_index.json')['segments']
        rows[kind] = {p.stem:read(p) for s in segments for p in (Path(s)/'task_rows').glob('*.json')}
    stats = '区间采用双侧95%题族整簇百分位bootstrap，重采样10,000次；同族题及其重复整体进入抽样。p值采用题族内净差整体换号的双侧符号置换，非零族数不超过18时枚举全部符号组合。只有净差来自不同题族的独立单配对时，数值才会与逐配对精确检验相同。'
    text = text.replace('Q明显改变选卡，', 'Q（经验卡从历史解题反馈积累的数值评分）改变了选卡，')
    text = text.replace('两组均未过门。多轮学习已经完成，本次结果不支持“反馈越多，正确率持续上升”。', '两组均未过门。本次比较学习后冻结的四份评分，没有检验轮次的单调趋势。')
    text = text.replace('五项探索经Holm校正后均未过门。多轮学习已经完成，本次结果不支持“反馈越多，正确率持续上升”。', '五项探索经Holm校正后均未过门。本次比较四份冻结评分，未进行单调趋势检验。')
    text = text.replace('本报告中的Joint表示同一份代码同时通过功能与安全测试。题次表示一道题的一次运行；两次重复仍是同一道题。两种模型覆盖的运行单元不同，结果按实际模型来源分别统计。',
        '表中Joint计数表示最终代码通过全部功能用例和全部安全用例的运行数。运行单元＝一道题×一个条件×一次重复；配对＝同题同重复的两个条件。案例中的“功能3/3”等另指测试用例数。两次重复不增加独立题数。')
    text = text.replace('本次授权的540次与820次均已交付。下列既有任务保持原状态；本次没有新增派发。',
        '本次授权的540次与820次均已交付。全量跳过来自用户限定的派发范围，与本次收益门分开；收益门用于解释结果。下列任务保持原状态。付费分类器评价指在298道审计题上新增模型生成调用；本次开发诊断只复用已完成结果，新增调用为0。')
    text = text.replace('57题检索与人工复核','57题经验卡适用性标注与人工复核').replace('新库、影响半径与Base/Plus','新经验库、代码改动范围与外部Base/Plus测试集')
    unit_note = ('经验卡是一条带适用条件与修改建议的可复用经验。候选卡是供模型选择的卡，选中卡是模型选入的卡，声明采用是模型说要使用；是否落实另看代码。blind、guard、hint、pitfall、cand是历史卡编号前缀，用于追溯。\n\n'
        'Q1是未使用学习Q的相似度对照；Q2_K0至Q2_K3是四份冻结评分条件。text-embedding-3-large编码查询和卡片，程序计算余弦相似度排列候选，DeepSeek作最终选卡。\n\n'
        '1017模型指deepseek-v4.1-flash-1017，Flash模型指deepseek-v4.1-flash。混合旧前缀表示同一运行复用1017成功响应后由Flash接续。Q全部来自此前1017学习；Flash组检验旧Q用于新接入模型的表现。来源标签覆盖该运行中的模型调用，分开统计。\n\n'
        +table(['名单','来自273题增长集','不同题族数'],[('Q回测','固定54题，约19.8%',54),('记忆标签','固定82题，约30.0%',80)])+
        '\n\n两份名单重合20题。名单在10月28日按固定种子冻结，选择未使用新成绩；规模来自两项独立授权。题族将相关题归为一组：Q的54题各自一族；标签82题含78个单题族与2个双题族。\n\n')
    text = text.replace('## Q冻结回测的完整结果\n\n', '## Q冻结回测的完整结果\n\n'+unit_note)
    text = text.replace('不据最佳轮次替换主要比较。', '不据最佳轮次替换主要比较。\n\n救回＝对照失败而目标通过；损失＝对照通过而目标失败；净变化＝救回减损失。'+stats+
        'Q的五项Holm校正在各模型来源内分别计算；原主要比较报告原p。本次另对两个纯来源的主比较补充Holm2，结果均为1；混合前缀仅作描述。')
    # 表中的分母和题族数直接来自已经冻结的统计。
    fixed=[];mode=None;active_cohort=None
    for line in text.splitlines():
        for k,n in NAMES.items():
            if line=='### '+n:active_cohort=k
        if line.startswith('| 执行来源 |'):mode='primary'
        elif line.startswith('| 来源 | 比较 |'):mode='qcomparisons'
        elif line.startswith('| 来源 | 条件 |'):mode='qarms'
        elif line.startswith('| 档位 | 保存 |'):mode='garms'
        elif line.startswith('| 来源 | 相对L1 |'):mode='gcomparisons'
        elif not line.startswith('|'):mode=None
        if line.startswith('|'):
            c=[x.strip() for x in line.strip('|').split('|')]
            if mode in ('primary','qcomparisons','gcomparisons'):
                pos=2 if mode=='primary' else 3
                if c[0] in ('执行来源','来源'):c.insert(pos,'题族数')
                elif re.fullmatch(r'[| :\-]+',line):c.insert(pos,'---')
                else:
                    cohort=next((k for k,n in NAMES.items() if n==c[0]),None)
                    if cohort:
                        data=q if mode!='gcomparisons' else g
                        comparison=('Q2_K3','Q1') if mode=='primary' else tuple(c[1].split('−')) if mode=='qcomparisons' else (c[1],'L1')
                        x=next(x for x in data['model_strata'][cohort]['comparisons'] if (x['target'],x['control'])==comparison)
                        if cohort=='mixed_saved_prefix':
                            if mode=='primary':c[4]=c[5]='仅描述'
                            else:c[-2]=c[-1]='仅描述'
                            if mode=='gcomparisons' and not x['measured_pairs']:c[3]=c[4]='无配对'
                        c.insert(pos,str(x['families']))
                line='| '+' | '.join(c)+' |'
            elif mode=='qarms' and c[0] in NAMES.values():
                cohort=next(k for k,n in NAMES.items() if n==c[0]);arm=q['model_strata'][cohort]['arms'][c[1]]
                c[4]=f"{arm['first']['joint']}/{arm['first']['measured']}" if arm['first']['measured'] else '未覆盖';line='| '+' | '.join(c)+' |'
            elif mode=='garms' and c[0].split(' ')[0] in LEVELS:
                level=c[0].split(' ')[0]
                # 这里由当前表的来源标题定位对应统计。
                cohort=active_cohort
                if cohort:arm=g['model_strata'][cohort]['arms'][level];c[3]=f"{arm['first']['joint']}/{arm['first']['measured']}" if arm['first']['measured'] else '未覆盖';line='| '+' | '.join(c)+' |'
        fixed.append(line)
    text='\n'.join(fixed)+'\n'
    recon=[]
    for cohort in ('1017_only','flash_only'):
        by={(str(r['task_id']),r['repeat'],r['condition']):r for r in rows['Q'].values() if r.get('model_cohort','1017_only')==cohort}
        pairs=[(r,by[str(r['task_id']),r['repeat'],'Q1']) for r in by.values() if r['condition']=='Q2_K3' and (str(r['task_id']),r['repeat'],'Q1') in by and r['final']['measured'] and by[str(r['task_id']),r['repeat'],'Q1']['final']['measured']]
        recon.append((NAMES[cohort],len(pairs),sum(a['final']['joint'] for a,b in pairs),sum(b['final']['joint'] for a,b in pairs)))
    text=text.replace('以下是各来源实际保存的条件结果；',table(['主比较来源','共同可测配对','配对内K3通过','配对内Q1通过'],recon)+
        '\n\n1017组K3的条件总数20/46含一个没有同来源Q1伙伴的成功单元；共同45对为19对17，净多2。条件表统计各自保存集合，净变化统计共同可测配对，不能直接相减两列总数。\n\n以下是各来源实际保存的条件结果；')
    text=text.replace('分母分别计算。','分母分别计算。“可测”指最终获得有效功能与安全判定；首次Joint按“通过数/首次可测数”显示，最终功能、安全和Joint的分母均为本行最终可测数。')
    text=text.replace('追加反馈扩大了评分覆盖，本次冻结回测没有确认额外轮次带来稳定Joint收益。',
        '四份冻结评分的反馈覆盖依次为445、610、702、769张，共1939张卡；累计反馈次数837、1483、2094、2837。本次回测没有更新Q。K3对K0未建立稳定配对收益，未进行单调趋势检验。\n\n'+
        table(['冻结评分','有反馈卡数／总卡数','累计反馈次数'],[(v['stage'],f"{v['feedback_covered']}/{v['cards']}",v['visits']) for v in read(Q/'analysis/q_coverage.json')]))
    text=text.replace('选卡变化已经发生；声明采用仍需核对代码。',
        '共同候选指同题同重复的查询、候选集合与分数一致。42/45、50/59是选中成员改变；33/45、36/59是在这些改选中声明采用新增成员，属于前者子集。代码落实的证据来自下面13个主要差异配对，未将全部声明自动当成落实。')
    text=text.replace('13个主要配对覆盖', '13个主要配对涉及11道不同题，1152与1218各出现两次。它们覆盖')
    text=text.replace('引入Q并检验选卡作用：已得到肯定回答。最终评分对正确率的作用：本次两个执行来源均未建立稳定提分。多轮反馈是否越来越有用：本次未建立单调或稳定收益。',
        'Q改变选卡：共同候选中1017组42/45、Flash组50/59的选中成员变化。K3相对相似度未过提分门；四份冻结评分的配对比较未显示稳定收益，未做单调趋势检验。')
    text=text.replace('另用原数据驱动验证器完成16次零API诊断，逐项结果与原记录一致。',
        '另对1152/1、1160/1、156/1、974/1、108/1、1028/1、683/1、1152/2这8个配对，各复算两组代码，共16次本地诊断，逐项与原记录一致。')
    text=text.replace('原功能失败、安全通过。', 'K3最终功能失败、安全通过。').replace('原safety明确要求','安全测试明确要求')
    text=text.replace('其代码已有html.escape，不能据此宣称存在XSS注入。','其代码已有html.escape，两组差异落实在模板内容。')
    text=text.replace('该配对差异包含输出类型契约，不能把所有功能损失统称为过度防御。','该配对的救回落实为除法结果保留float返回类型。')
    text=text.replace('新增卡的采用没有避免模型实现错误；该损失不能解释为算术安全原理无效。','新增卡被采用后仍出现AST根节点漏存结果的实现错误。')
    text=text.replace('安全组通过还可能来自一概拒绝，不能单独作为安全能力进步。','该代码一概拒绝，安全组通过与合法请求全部失败同时发生。')
    text=text.replace('自动生成的73条探索案例队列不等于全部已审，本阶段完成的是全部13个主要救回与损失配对。','73条自动探索候选中，已完成全部13个主要救回与损失配对的代码复盘；其余候选未统一审阅，不用于签署原因。')
    case_lines=[]
    for line in text.splitlines():
        case_lines.append(line)
        match=re.match(r'^#### \[(\d+)／(\d+)\]',line)
        if match:
            tid,rep=match.groups();verdict=[]
            for condition,name in [('Q1','相似度组'),('Q2_K3','K3组')]:
                f=rows['Q'][f'{tid}_r{rep}_{condition}']['final']
                verdict.append(name+'：'+('未测' if not f['measured'] else '、'.join(k+('通过' if f[field] else '失败') for k,field in [('功能','functional'),('安全','security'),('Joint','joint')])))
            case_lines += ['', '最终判定：'+'；'.join(verdict)+'。']
    text='\n'.join(case_lines)+'\n'
    text=text.replace('L1只给所选卡的风险提醒；','L1只给所选卡的风险提醒，每题最多三张；固定L1表示固定档位，提醒内容随题目检索卡变化；')
    text=text.replace('通过数只来自已经完成原验证的单元。','“首次Joint”按“通过数/首次可测数”显示；最终通过数的分母是本行最终可测数。')
    text=text.replace('四项比较按真实题族整体重采样并进行Holm4校正。','四项比较在每个模型来源内分别做Holm4校正。'+stats)
    text=text.replace('当前结果按观察差异解释，没有确认更长或更短记忆稳定占优。','当前结果没有确认某个呈现变体稳定占优；L3与L4分别比较，不作逐级加长解释。')
    text=text.replace('更长的来源示例档','L4来源示例档').replace('两个较长记忆档损失配对','L3与L4各一个损失配对')
    text=text.replace('模板中只绑定headers/body或通用操作，未形成正确的媒体类型限制。','L3的操作模板只标出headers/body等目标元素或通用操作，没有明确的媒体类型白名单。')
    text=text.replace('来源示例在权限修改题救回，也在资料访问题损失；目标模板仍可能遗漏媒体类型规则。','L4在1023权限修改题救回，在1012资料访问题损失；L3在497题没有补齐媒体类型规则。')
    text=text.replace('五档均无2/2成功档，该题保持未知标签，不指定L0；增加记忆长度未在本题消除该失败。',
        '110仍计入82题表现；没有档连续两次通过，因此按原规则不给首选标签。隐藏文件要求未在题面明示是单独的契约问题，没有据此剔题、改分或改标签规则。')
    text=text.replace('换模后32题获得同一新模型的完整十格结果，其中训练11题、开发21题。',
        '十格＝同题五档×两次重复。原训练61题中11题取得纯Flash十格，其余50题跨模型来源；原开发21题全部取得纯Flash十格，划分未改变。32题来自原顺序执行中换模后的完成集合，未按成绩重抽；本次训练判断限定这11题。')
    text=text.replace('标签只在十格均可测、有卡片呈现、至少一个档两次Joint通过且输入成本已知时成立；',
        '标签要求同题十格均可测、至少一个带卡条件曾呈现卡片、至少一档两次均Joint通过且成功档输入成本已知；首选档允许L0。')
    text=text.replace('P7输入包含该档分析、规划、代码生成与修复请求，公共初析、检索和材料制作另外计费。',
        'P7是本实验的分析—规划—生成流程。成本标签使用该档成功返回请求已报告的输入Token，含修复；公共初析、检索、材料制作和失败请求未知用量单独列账。它用于比较已知输入量，不是完整货币支出。')
    text=text.replace('各档两次结果完全相同','五档的两次结果序列相同').replace('两次均体现同向档位差异','一档两次均过，另一档两次均不过')
    text=text.replace('逐题上限使用同一批已经看到答案的结果，只用于判断选档空间。',
        '例如五档都为[过,不过]，归入结果序列相同；五档序列不同但没有“一档均过、另一档均不过”，归入伴随重复波动。最佳固定档为L4，分母22＝11道训练题×2次；固定档和逐题上限都在这11题上事后选取，只用于训练侧选档空间诊断。')
    text=text.replace('| 分类器不能是常量 | multinomial_logistic |','| 分类器不能是常量 | 多项逻辑回归；开发实际选择L0、L1、L4三档 |')
    text=text.replace('L0=2族；L4=1族；L1=2族','总共5族：L0=2、L4=1、L1=2；没有一类达到5族')
    text=text.replace('下面保留该候选的开发诊断。工程门未过，开发结果不授予正式策略资格。',
        '已拟合的是离线诊断候选，正式分类器没有冻结或上线。开发42次＝21道题×2次，仅复用已完成代码；训练侧最佳固定档为L4，L1是预定参照档。')
    text=text.replace('大写外部名称数','大写标识符数')
    model=read(G/'analysis/candidate_model.json')
    text=text.replace('避免同族题同时进入拟合与验证。','避免同族题同时进入拟合与验证。完整训练族三折分别3、5、3族；其中有标签可拟合的族分别1、3、1。C=0.1、1、10的验证族平均Joint均为0.54545，输入均值也并列，按规则选C=0.1。'+
        '最高类别概率不足0.5时使用训练侧固定L4。简单规则在无卡时选L0；题面有三字符以上全大写标识或副作用词时选L3，否则L2。开发比较未参与调参。')
    for a,b in [('classifier−fixed_L1','离线候选−固定L1'),('classifier−fixed','离线候选−固定L4'),('classifier−rule','离线候选−简单规则'),('| fixed |','| 固定L4 |'),('| rule |','| 简单规则 |'),('| classifier |','| 离线候选 |'),('| fixed_L1 |','| 固定L1 |')]:text=text.replace(a,b)
    text=text.replace('[训练资格原件]', '开发差值的区间同样为双侧题族bootstrap；没有任何救回时，重采样净差只能为负或零，因此上端可等于零。\n\n[训练资格原件]',1)
    preferred=[]
    labels=read(G/'analysis/flash_only_labels.json')
    for level in LEVELS:
        selected=[x for x in labels if x['preferred_level']==level];group=[rows['G'][f"{x['task_id']}_r{rep}_{level}"] for x in selected for rep in (1,2)]
        adopted=sum(any(a['generation'].get('cards') for a in r['attempts']) for r in group)
        presented=sum(any(a['generation'].get('retrieved_cards') for a in r['attempts']) for r in group)
        preferred.append((level,len(selected),len(group),adopted,presented-adopted if level!='L0' else 0,len(group) if level=='L0' else 0))
    text=text.replace('标签表示本次观察的成本首选，不能等同于该题永久需要的记忆程度。','标签表示本次“通过前提下已知输入最少”的呈现，不要求声明采用卡；下面明确列出首选档的采用状态。\n\n'+
        table(['首选档','题数','两次运行数','至少一次声明采用','有呈现但全程未采用','L0无卡'],preferred))
    text=text.replace('540次一致、差异0、新HTTP为0。','535次有效判定与5次未测状态均重现，差异0，新HTTP为0。')
    text=text.replace('两批共1360次交付，6次最终结果未测。','原生回放按原流程消费保存响应并重验代码，截断等未测状态也需一致；Q为535个有效判定加5个未测状态，标签为819加1。回放通过不把未测转为成功。\n\n两批共1360次交付，6次最终结果未测。')
    missing=read(Q/'analysis/unmeasured_review.json')['records']
    text=text.replace('[未测原件清单]',table(['Q未测运行单元','首次状态','最终未测原因'],[(v['cell_id'],
        '未测' if not rows['Q'][v['cell_id']]['attempts'][0].get('measured') else '功能'+('通过' if rows['Q'][v['cell_id']]['attempts'][0].get('functional') else '失败')+'／安全'+('通过' if rows['Q'][v['cell_id']]['attempts'][0].get('security') else '失败'),
        v['error_type'] or '原验证记录未测，未报告具体错误') for v in missing])+
        '\n\n新模型683题第2次的Q1最终Joint失败。若缺失K3按失败或成功两种极端补值，主配对净差分别为−1/60、0/60，均不为正；正式结果仍保持59对的−1。截断与输出长度相关，不假定随机缺测，也不改写原值。\n\n[未测原件清单]',1)
    text=text.replace('原累计上限8100万Token、21600次HTTP及8640次Embedding HTTP保持不变。','Q这540次的原累计上限为8100万Token、21600次HTTP，其中Embedding至多8640次。')
    text=text.replace('统计来自物理账本，复用行中的usage没有重复相加。原预算累计保留；',
        '实际网络请求共10,078次，未知用量58次占0.58%；缓存复用不增加请求。这里报告Token与请求数，没有换算货币。已报告Token不包含未知部分；')
    text=text.replace('L4不包含L3，五档不是逐级累计。','L3与L4分别在L2上扩展，为两个平行变体。')
    text=text.replace('它用于比较已知输入量，不是完整货币支出。','该指标计量已知输入Token；货币支出未换算。')
    text=text.replace('“首次Joint”按“通过数/首次可测数”显示；最终通过数的分母是本行最终可测数。','“首次Joint”按“通过数/首次可测数”显示；最终通过数的分母是本行最终可测数。保存0表示该来源无此档记录。')
    text=text.replace('未填回正式请求。','独立于两批账本，未填回正式请求。\n\n'+table(['固定批次预算','Token上限','已知Token','按已知用量计算的账面差额','HTTP上限／已用'],[
        ('Q回测','81,000,000','35,806,162','45,193,838','21,600／4,328'),('记忆标签','221,400,000','32,039,253','189,360,747','40,672／5,750')])+ '\n\n标签Embedding上限13,120次，已用480次；Q已用426次。账面差额未扣除未知用量，不作为可继续花费额度。本批没有追加派发。')
    text=text.replace('再由原模型选卡','再由对应生成模型选卡').replace('原模型来源','1017模型来源').replace('原模型5次救回','1017模型5次救回').replace('原模型和混合前缀','1017模型和混合前缀').replace('无1017模型','Flash模型').replace('无1017来源','Flash来源')
    support={'primary_reconciliation':recon,'preferred_level_adoption':preferred,'primary_Holm2':{'1017_only':1.0,'flash_only':1.0},'unknown_HTTP_fraction':58/10078,'original_native_results_changed':False}
    (OUT/'review/readability_support.json').write_text(json.dumps(support,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    return text


def report():
    q, g = read(Q/'analysis/summary.json'), read(G/'analysis/summary.json')
    assert q['status'] == g['status'] == 'completed'
    assert read(Q/'replay_result.json')['completed'] == 540
    assert read(G/'replay_result.json')['completed'] == 820
    review = read(G/'analysis/agent_case_review.json')
    labels = read(G/'analysis/flash_only_labels.json')
    opportunity = read(G/'analysis/training_opportunity.json')
    repeats = read(G/'analysis/repeat_stability.json')
    ready = g['training_readiness']
    queue = read(G/'analysis/case_review_queue.json')
    checked = {tuple(v['label_key']) for v in review['records'] if 'label_key' in v}
    assert all(tuple(v['label_key']) in checked for v in queue)
    pure_labels = [v for v in labels if v['label_status'] != 'incomplete_grid']
    pref = Counter(v['preferred_level'] for v in pure_labels if v['preferred_level'])
    check = ready['readiness']
    lines = ['# 11月1日Q回测与记忆标签结果分析', '<a id="main"></a>', '',
        '本次完成老师提出的经验评分与记忆程度两个问题的固定小批验证：Q回测54题540次，记忆标签82题820次，均已完整原生回放。Q明显改变选卡，稳定提分门未通过；记忆程度的收益门也未通过，分类器训练资格未达标。下一步保留这些结论，排队全量继续跳过。', '',
        '<a id="r1"></a>', '## Q改变选卡后提高正确率了吗？', '',
        '最终K3相对相似度选卡，原模型来源净多2/45次，无1017模型来源净少1/59次，两组均未达到预登记的提分门槛。共同候选中改选成员分别42/45、50/59题次；评分已影响选卡，正确率改善没有成立。'+
        '[【A1】完整比较](#q_a1) · [【A2】13个主要差异配对](#q_a2)', '',
        '<a id="r2"></a>', '## 增加反馈轮次持续提高正确率了吗？', '',
        'K3相对K0，原模型来源净少1/46次，无1017模型来源净多1/58次；五项探索经Holm校正后均未过门。多轮学习已经完成，本次结果不支持“反馈越多，正确率持续上升”。'+
        '[【A1】轮次和重复](#q_a1)', '',
        '<a id="r3"></a>', '## 按题选择记忆程度具备训练和提分依据了吗？', '',
        f'新模型来源中，L0/L1/L2/L3/L4最终Joint分别为58/113、58/114、56/114、59/114、59/114。完整同模型标签覆盖32题，其中训练11题、开发21题；训练题事后逐题选档上限与最佳固定档同为13/22次。可用训练标签只有{check["training_label_families"]}个题族，未达到30族门槛，本次不发布分类器。'+
        '[【A3】五档比较](#g_results) · [【A4】代码差异](#g_cases) · [【A5】训练判断](#g_training)', '',
        '本报告中的Joint表示同一份代码同时通过功能与安全测试。题次表示一道题的一次运行；两次重复仍是同一道题。两种模型覆盖的运行单元不同，结果按实际模型来源分别统计。', '',
        '<a id="todo"></a>', '## 哪些任务保持未执行？', '',
        '本次授权的540次与820次均已交付。下列既有任务保持原状态；本次没有新增派发。', '',
        table(['任务','原因','下一步与当前状态'], [
            ('Q跨题族与排队全量','本次仅授权固定54题回测','保留冻结Q及已完成结果；全量继续跳过。'),
            ('追加82题标签与付费分类器评价','训练资格未过；没有追加授权','保存候选诊断，不追加题、不冻结正式分类器、不启动付费评价。'),
            ('57题检索与人工复核','该批不在本次两批续跑范围；人工确认仍缺失','保留已有候选与初标，待独立安排续段及盲审；本报告不授予人工结论。'),
            ('新库、影响半径与Base/Plus','本次没有执行授权或原环境依赖未齐','保留现有计划、结果与依赖状态，不启动。')]), '',
        '[返回研究问题](#main)', '', '<a id="q_top"></a>', '## Q冻结回测的完整结果', '',
        '四份评分来自此前1017模型的四轮学习；换模后继续冻结使用，未用新模型重新学习。', '']
    # 复用已交付的完整Q数据；将长案例表改成逐例段落，保留全部文字与链接。
    source = (ROOT/'docs/research/reports/CLE-RPT-20261101-Q-retest-v01/11月1日Q冻结回测结果与分析.md').read_text(encoding='utf-8')
    source = source[source.index('<a id="a1"></a>'):]
    source = re.sub(r'<a id="([^"]+)"></a>', lambda m:'<a id="q_'+m[1]+'"></a>', source)
    source = source.replace('](#', '](#q_')
    source = source.replace('记忆标签按原82题范围接续，完成后再独立评估标签与训练资格。', '记忆标签的完整结果与训练判断见本报告【A3】至【A5】。')
    converted=[]
    for line in source.splitlines():
        if line.startswith('| 题目／重复 |') or (line.startswith('|---|---|---|---|') and converted and converted[-1]=='CASE_TABLE'):
            converted.append('CASE_TABLE'); continue
        if line.startswith('| [') and '](' in line and line.count('|')==5:
            cells=[v.strip() for v in line.strip('|').split('|')]
            converted.extend(['', '#### '+cells[0]+' '+cells[1], '', cells[2], '', cells[3], '']); continue
        if line.startswith('#'):line='#'+line
        converted.append(line)
    lines.extend(v for v in converted if v!='CASE_TABLE')
    lines += ['', '<a id="g_results"></a>', '## 【A3】五种记忆程度改变了正确率吗？', '',
        '**结论：** 四个替代档相对风险提醒的收益门均未通过。新模型来源中，目标模板档与来源示例档各多通过1次；配对比较同时包含救回与损失，题族区间跨零。', '',
        '老师要求按题选择记忆程度。先保持任务、卡片来源与模型流程一致，比较五种呈现，再用训练题判断能否学会选档。'+
        'L0全程无卡；L1只给所选卡的风险提醒；L2给原则与适用条件；L3在L2上加入针对当前题的操作模板；L4在L2上加入来源题及代码示例。L4不包含L3，五档不是逐级累计。', '',
        '1. 固定82题、五档、两次独立代码重复，共820个运行单元；训练61题、开发21题的原题族划分保持。',
        '2. 同题同次重复共享首次初析和首次卡片来源，再按档位变换呈现；修复各档独立，最多三卡、一次修复，Q和卡片正文不更新。',
        '3. 保留换模前247次和成功请求前缀，其余使用无1017模型；旧模型、新模型与混合前缀分开统计。',
        '4. 820次交付后完整原生零API回放；统计同来源配对、重复波动、标签和训练资格，再阅读主要案例。', '',
        '下面按执行来源列出全部首次与最终结果。保存数包括未测记录；通过数只来自已经完成原验证的单元。', '']
    for cohort, data in g['model_strata'].items():
        lines += ['### '+NAMES[cohort], '', table(['档位','保存','可测','首次Joint','最终功能','最终安全','最终Joint'], [
            (level+' '+LEVELS[level],arm['final']['completed'],arm['final']['measured'],arm['first']['joint'],arm['final']['functional'],arm['final']['security'],arm['final']['joint'])
            for level,arm in data['arms'].items()]), '']
    lines += ['### 同来源配对与重复', '', table(['来源','相对L1','可测配对','救回／损失','净变化','题族95%区间 百分点','Holm4 p'], [
        (NAMES[c],x['target'],x['measured_pairs'],f'{x["wins"]}／{x["losses"]}',f'{x["wins"]-x["losses"]:+d}',
         '无配对' if x['ci95_family_bootstrap'] is None else '至'.join(map(pp,x['ci95_family_bootstrap'])),x.get('exploratory_Holm4_p','无配对'))
        for c,d in g['model_strata'].items() for x in d['comparisons']]), '',
        '四项比较按真实题族整体重采样并进行Holm4校正。正向门槛为净差为正、95%区间下界大于零、校正p小于0.05；负向门槛取相反方向。当前结果按观察差异解释，没有确认更长或更短记忆稳定占优。', '',
        table(['新模型档位','两次均可测题数','Joint结果翻转题数','两次均通过','两次均失败'],[(l,d['two_measured_repeat_tasks'],d['Joint_flip_tasks'],d['both_pass'],d['both_fail']) for l,d in repeats['flash_only'].items()]), '',
        '原模型和混合前缀均没有同来源内的完整双重复题，重复波动保留未覆盖状态。新模型五档各有32题可比较两次重复，4至6题的Joint发生翻转，说明单次成功不足以给题目指定可靠档位。', '',
        f'[完整统计]({G.as_posix()}/analysis/summary.json) · [重复记录]({G.as_posix()}/analysis/repeat_stability.json) · [返回记忆问题](#r3)', '',
        '<a id="g_cases"></a>', '## 【A4】代码揭示了什么具体原因？', '',
        '实际阅读覆盖四类完整标签，以及两个较长记忆档损失配对，共6道题。14次原验证器本地诊断与正式记录一致，新增HTTP为0。', '']
    for item in review['records']:
        lines += ['### '+item['task_id']+' '+item.get('comparison', '标签案例'), '', item['result'], '', item['observed_code'], '', item['card_and_plan'], '', item['interpretation'], '',
            '证据：'+' · '.join(f'[{Path(p).stem}]({Path(p).as_posix()})' for p in item['source_rows']), '']
    lines += ['外部配置是否保留、计划是否落实、核心安全规则是否写对，直接区分了这些案例的结果。来源示例在权限修改题救回，也在资料访问题损失；目标模板仍可能遗漏媒体类型规则。上述配对没有单卡置换，卡片展示、采用判断、规划与生成分别记述。', '',
        f'[实际案例清单]({G.as_posix()}/analysis/agent_case_review.json) · [14次诊断]({G.as_posix()}/analysis/agent_case_diagnostic_replay.json) · [返回记忆问题](#r3)', '',
        '<a id="g_training"></a>', '## 【A5】这些标签支持训练分类器吗？', '',
        '**结论：** 本次训练资格未通过，不发布正式分类器，不启动付费策略评价。', '',
        f'换模后32题获得同一新模型的完整十格结果，其中训练11题、开发21题。其余50题含不同模型来源，继续保留原逐格结果，未拼成训练标签。32题中两次都成功档的成本首选为：'+
        '、'.join(f'{l} {pref.get(l,0)}题' for l in LEVELS)+f'；其余{sum(v["label_status"]=="no_stable_success" for v in pure_labels)}题没有任何档连续两次通过。', '',
        '标签只在十格均可测、有卡片呈现、至少一个档两次Joint通过且输入成本已知时成立；在成功档中选平均P7输入Token最低者。P7输入包含该档分析、规划、代码生成与修复请求，公共初析、检索和材料制作另外计费。标签表示本次观察的成本首选，不能等同于该题永久需要的记忆程度。', '',
        table(['训练11题的观察','数量'], [
            ('各档两次结果完全相同',opportunity['task_counts'].get('no_observed_correctness_difference',0)),
            ('档位差异伴随重复波动',opportunity['task_counts'].get('difference_with_repeat_variation',0)),
            ('两次均体现同向档位差异',opportunity['task_counts'].get('same_direction_difference_in_both_repeats',0)),
            ('最佳固定档Joint',f'{opportunity["best_observed_fixed_joint"]}/{opportunity["paired_observations"]}'),
            ('事后逐题选档Joint上限',f'{opportunity["hindsight_per_task_joint"]}/{opportunity["paired_observations"]}'),
            ('事后逐题选档额外通过数',opportunity['potential_net_joint'])]), '',
        '逐题上限使用同一批已经看到答案的结果，只用于判断选档空间。它与最佳固定档相同，本次训练题没有显示额外正确率空间。', '',
        table(['原工程门槛','实际'],[
            ('至少30个可用训练标签题族',check['training_label_families']),
            ('至少两类各有5个训练题族','；'.join(f'{k}={v}族' for k,v in check['class_training_families'].items())),
            ('分类器不能是常量',read(G/'analysis/candidate_model.json')['type']),
            ('开发路由至少使用两档','、'.join(check['development_routed_levels']))]), '',
        '离线候选沿用原8个题目特征和逻辑回归，不读取Q、CWE、题号或未来修复卡。正则参数仅在训练题族三折上按Joint选取，再看21道开发题；下面保留该候选的开发诊断。工程门未过，开发结果不授予正式策略资格。', '',
        '8个特征分别为：公开题面长度、参数说明词数、上下文长度、大写外部名称数、是否说明返回值、是否说明异常、副作用操作词数、安全要求长度。模型输入只有题面；输出一个记忆档位。训练侧将同题族放在同一折，避免同族题同时进入拟合与验证。', '',
        table(['开发策略','题次','功能通过','安全通过','Joint通过','实际档位'],[
            (policy,len(items),sum(r['final']['functional'] for r in items),sum(r['final']['security'] for r in items),sum(r['final']['joint'] for r in items),
             '；'.join(f'{k}={n}' for k,n in sorted(Counter(r['selected_level'] for r in items).items())))
            for policy,items in read(G/'analysis/development_routes.json').items()]), '',
        table(['开发比较','可测配对','救回／损失','净变化','题族95%区间 百分点'],[
            (x['target']+'−'+x['control'],x['measured_pairs'],f'{x["wins"]}／{x["losses"]}',f'{x["wins"]-x["losses"]:+d}', '无配对' if x['ci95_family_bootstrap'] is None else '至'.join(map(pp,x['ci95_family_bootstrap'])))
            for x in ready['development_comparisons']]), '',
        f'[训练资格原件]({G.as_posix()}/analysis/training_readiness.json) · [训练题选档空间]({G.as_posix()}/analysis/training_opportunity.json) · [完整标签]({G.as_posix()}/analysis/flash_only_labels.json) · [返回记忆问题](#r3)', '',
        '<a id="g_cost"></a>', '## 【A6】交付范围和费用是什么？', '',
        table(['实验','独立题数','完成／计划','可测','未测','原生回放'],[('Q冻结回测',54,'540/540',535,5,'540/540一致'),('五档记忆标签',82,'820/820',819,1,'820/820一致')]), '',
        '两批共1360次交付，6次最终结果未测。Q的5次原因保留在前文；标签863题第2次L0修复输出截断，最终没有代码可测，原首次代码与失败记录保留。未测单列，不当作已验证失败，也不按成绩重新生成。', '',
        table(['实验','实际请求模型','物理HTTP','已报告Token','未知用量HTTP'],[
            (package,model,d['physical_HTTP'],f'{d["known_reported_tokens"]:,}',d['unknown_or_inflight_HTTP'])
            for package,data in [('Q',q),('记忆标签',g)] for model,d in data['physical_cost_by_requested_model'].items()]), '',
        f'两批累计已报告{sum(d["known_reported_tokens"] for data in (q,g) for d in data["physical_cost_by_requested_model"].values()):,} Token，'+
        f'{sum(d["unknown_or_inflight_HTTP"] for data in (q,g) for d in data["physical_cost_by_requested_model"].values())}次HTTP用量未知。统计来自物理账本，复用行中的usage没有重复相加。原预算累计保留；独立长选卡控制测试26,527 Token另行登记，未填回正式请求。', '',
        f'[Q回放证书]({Q.as_posix()}/replay_result.json) · [标签回放证书]({G.as_posix()}/replay_result.json) · [标签未测记录]({G.as_posix()}/analysis/unmeasured_review.json)', '',
        '这两批固定实验、原生回放和分析已完成。Q的主要作用已确认到选卡，稳定提分未成立；记忆分档在具体题上出现救回与损失，训练资格与整体提分门均未过。本批付费派发结束，排队全量和其他待做保持原状态。', '',
        '[返回研究问题](#main) · [当前实验计划](D:/Agent/sct-security-generation/docs/research/CLE-PLAN-20261101-to-20261104-v01.md)']
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT/'review').mkdir(exist_ok=True)
    MD.write_text(clarify('\n'.join(lines)+'\n',q,g), encoding='utf-8')
    print(MD)


def word():
    from tools.build_cle_weekly_package_20261028 import build_word
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    path=build_word(MD,'report');doc=Document(path)
    for table_obj in doc.tables:
        headers=[c.text for c in table_obj.rows[0].cells];n=len(headers)
        if n==8:
            widths=([2.1,2.3,1.1,1.1,1.8,1.8,1.8,1.8] if '条件' in headers else
                    [1.8,2.6,1.2,1.0,1.3,1.1,2.9,1.4] if '比较' in headers else
                    [1.8,1.3,1.0,1.3,1.2,2.9,1.5,2.3] if '判断' in headers else
                    [1.8,1.2,1.3,1.0,1.4,1.2,3.0,1.5])
        elif n==7:widths=[3.4,1.2,1.2,2.0,1.7,1.7,1.7]
        elif n==6:widths=[3.1,1.3,2.0,2.0,2.3,4.2]
        elif n==5:widths=([2.0,4.8,2.5,3.1,2.7] if '实际请求模型' in headers else
                          [2.2,3.1,3.0,4.1,4.6] if 'Token上限' in headers else [3.4,1.8,2.5,2.4,4.8])
        elif n==4:widths=[4.5,2.4,3.5,3.5]
        elif n==3:widths=[3.8,4.5,7.7]
        else:widths=[10,5]
        widths=[17*x/sum(widths) for x in widths];table_obj.autofit=False
        pr=table_obj._tbl.tblPr;borders=pr.find(qn('w:tblBorders'))
        if borders is None:borders=OxmlElement('w:tblBorders');pr.append(borders)
        for edge in ('top','left','bottom','right','insideH','insideV'):
            node=OxmlElement('w:'+edge);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9');borders.append(node)
        for k,width in enumerate(widths):table_obj.columns[k].width=Cm(width)
        for i,row in enumerate(table_obj.rows):
            for k,cell in enumerate(row.cells):
                cell.width=Cm(widths[k]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing=1.1
                    paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER if k>0 and len(cell.text)<22 else WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:run.font.size=Pt(9.5)
    # 覆盖超链接内部run，统一汉字与拉丁字符字体，保留标题字号和强调。
    for run in doc.element.iter(qn('w:r')):
        rpr=run.find(qn('w:rPr'))
        if rpr is None:rpr=OxmlElement('w:rPr');run.insert(0,rpr)
        fonts=rpr.find(qn('w:rFonts'))
        if fonts is None:fonts=OxmlElement('w:rFonts');rpr.append(fonts)
        for key,value in (('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','宋体')):fonts.set(qn('w:'+key),value)
    doc.save(path);return path


if __name__ == '__main__':
    report()
