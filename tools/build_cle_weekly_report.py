"""Build linked weekly Word reports from the registered Markdown sources."""
from pathlib import Path
import argparse
import json
import re
from urllib.parse import urlparse

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / 'docs/research/reports/CLE-RPT-20261022-20261028-weekly-v04'
TEMPLATE = ROOT / 'docs/research/templates/CLE-TPL-weekly-experiment-report-v03.md'
SOURCE = ROOT / 'results/curated/cle_v41_completion_20261026/summary.json'
BLACK = RGBColor(0, 0, 0)


def figures():
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    data = json.loads(SOURCE.read_text(encoding='utf-8-sig'))
    stages = {s['id']: s for s in data['stages']}
    out = REPORT / 'figures'
    out.mkdir(exist_ok=True)
    chinese = Path('C:/Windows/Fonts/simsun.ttc')
    latin = Path('C:/Windows/Fonts/times.ttf')
    font_manager.fontManager.addfont(str(chinese))
    font_manager.fontManager.addfont(str(latin))
    plt.rcParams.update({'font.family': [font_manager.FontProperties(fname=latin).get_name(),
                                       font_manager.FontProperties(fname=chinese).get_name()],
                         'font.size': 11, 'axes.unicode_minus': False, 'text.color': 'black',
                         'axes.labelcolor': 'black', 'xtick.color': 'black', 'ytick.color': 'black'})

    def axes_style(ax, ymax=100):
        ax.set_ylim(0, ymax)
        ax.spines[['top', 'right']].set_visible(False)
        ax.spines[['bottom', 'left']].set_color('#9CA3AF')
        ax.yaxis.grid(True, color='#E5E7EB', linewidth=.6)
        ax.set_axisbelow(True)
        ax.set_ylabel('通过率（%）')

    def joint(stage, arms, names, file):
        fig, ax = plt.subplots(figsize=(7.1, 2.85))
        vals = [stages[stage]['arms'][a] for a in arms]
        bars = ax.bar(names, [100*v['joint']/v['completed'] for v in vals],
                      color=['#64748B' if i == 0 else '#648BA8' for i in range(len(vals))], width=.57)
        axes_style(ax)
        for bar, v in zip(bars, vals):
            y = bar.get_height()
            ax.text(bar.get_x()+bar.get_width()/2, y+2, f"{v['joint']}/{v['completed']}\n{y:.2f}%", ha='center', fontsize=10)
        fig.tight_layout(pad=.7)
        fig.savefig(out/file, dpi=210, facecolor='white')
        plt.close(fig)

    joint('D03b',['Q0','Q1','Q2','Q3'],['Q0\n原CLE','Q1\n相似度','Q2\n学习Q','Q3\n打乱Q'],'a1_q_joint.png')
    joint('B02',['B0','BM','BF','BS','BFS'],['无记忆','混合卡','功能视图','安全视图','分区组合'],'a2_fs_joint.png')
    joint('C02',['L0','L1','L2','L3','L4'],['L0\n无记忆','L1\n风险提醒','L2\n机制前提','L3\n操作模板','L4\n来源示例'],'a3_levels_joint.png')
    joint('C04b',['classifier','fixed','rule'],['分类器','固定L1','公开规则'],'a4_classifier_joint.png')

    fig, ax = plt.subplots(figsize=(7.1, 2.8))
    for i,(arm,label,color) in enumerate([('P7','原P7','#64748B'),('PUBLIC','公开检查','#648BA8')]):
        d = stages['D02']['arms'][arm]
        x = [j+(i-.5)*.34 for j in range(3)]
        bars = ax.bar(x,[d[k]/96*100 for k in ['functional','security','joint']],width=.32,color=color,label=label)
        for b,k in zip(bars,['functional','security','joint']):
            ax.text(b.get_x()+b.get_width()/2,b.get_height()+2,f"{d[k]}/96",ha='center',fontsize=10)
    axes_style(ax)
    ax.set_xticks(range(3),['功能F','安全S','联合Joint'])
    ax.legend(frameon=False,loc='upper right',ncol=2)
    fig.tight_layout(pad=.7)
    fig.savefig(out/'a5_public_metrics.png',dpi=210,facecolor='white'); plt.close(fig)

    fig, ax = plt.subplots(figsize=(7.1, 2.8))
    groups = [('S2_M1','protect_none','protect_m1'),('S4','CONTROL','SELECTED')]
    for i,(label,color) in enumerate([('无经验','#64748B'),('来源经验M1','#648BA8')]):
        vals = [stages[stage]['arms'][keys[i]] for stage,*keys in groups]
        bars = ax.bar([j+(i-.5)*.34 for j in range(2)], [v['joint']/v['completed']*100 for v in vals],width=.32,color=color,label=label)
        for b,v in zip(bars,vals):
            ax.text(b.get_x()+b.get_width()/2,b.get_height()+2,f"{v['joint']}/{v['completed']}",ha='center',fontsize=10)
    axes_style(ax)
    ax.set_xticks([0,1],['开发阶段（12题族）','保留确认（24题族）'])
    ax.legend(frameon=False,loc='upper right',ncol=2)
    fig.tight_layout(pad=.7)
    fig.savefig(out/'a6_m1_comparison.png',dpi=210,facecolor='white'); plt.close(fig)

    ids=['B03R','B02','S2_M1','S2_P13_P22','D02','C02','C04b','D03b','B05','S4','D03a']
    labels=['关系提醒','功能安全','M1开发','生成模板','公开检查','五档记忆','分类器确认','Q代码对照','经验编辑','M1保留确认','仅选卡']
    fig,ax=plt.subplots(figsize=(7.1,3.7))
    values=[stages[i]['completed'] for i in ids]
    bars=ax.barh(labels,values,color=['#648BA8']*10+['#B79C73'],height=.64)
    ax.invert_yaxis(); ax.set_xlim(0,550)
    ax.spines[['top','right']].set_visible(False)
    ax.set_xlabel('完成的实验单元数（含条件、重复及复用）')
    ax.xaxis.grid(True,color='#E5E7EB',linewidth=.6); ax.set_axisbelow(True)
    for b,v in zip(bars,values): ax.text(v+6,b.get_y()+b.get_height()/2,str(v),va='center',fontsize=10)
    fig.tight_layout(pad=.7)
    fig.savefig(out/'a7_completion.png',dpi=210,facecolor='white'); plt.close(fig)


def fonts(style, size, east='宋体', bold=False):
    style.font.name='Times New Roman'
    style.font.size=Pt(size)
    style.font.color.rgb=BLACK
    style.font.bold=bold
    rfonts=style.element.get_or_add_rPr().rFonts
    for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        rfonts.attrib.pop(qn('w:'+attr),None)
    rfonts.set(qn('w:eastAsia'),east)


def add_hyperlink(paragraph,label,target,base):
    h=OxmlElement('w:hyperlink')
    if target.startswith('#'):
        h.set(qn('w:anchor'),target[1:])
    else:
        if not urlparse(target).scheme:
            target=(base/target).resolve().as_uri()
        h.set(qn('r:id'),paragraph.part.relate_to(target,RT.HYPERLINK,is_external=True))
    h.set(qn('w:history'),'1')
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    col=OxmlElement('w:color'); col.set(qn('w:val'),'000000'); pr.append(col)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u)
    rf=OxmlElement('w:rFonts');rf.set(qn('w:eastAsia'),'宋体');rf.set(qn('w:ascii'),'Times New Roman');rf.set(qn('w:hAnsi'),'Times New Roman');pr.append(rf)
    r.append(pr); t=OxmlElement('w:t'); t.text=label;r.append(t);h.append(r);paragraph._p.append(h)


def inline(p,text,base):
    pattern=r'(\[[^\]]+\]\([^)]+\)|\*\*.+?\*\*|`[^`]+`)'
    for part in re.split(pattern,text):
        if not part: continue
        m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',part)
        if m: add_hyperlink(p,m[1],m[2],base)
        elif part.startswith('**'): p.add_run(part[2:-2]).bold=True
        elif part.startswith('`'): p.add_run(part[1:-1])
        else: p.add_run(part)


def bookmark(p,name,number):
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(number));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(number))
    p._p.insert(0,start);p._p.append(end)


def new_numbering(doc):
    numbering=doc.part.numbering_part.element
    abstract_id=max([int(x.get(qn('w:abstractNumId'))) for x in numbering.findall(qn('w:abstractNum'))]+[-1])+1
    aid=OxmlElement('w:abstractNum'); aid.set(qn('w:abstractNumId'),str(abstract_id))
    lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),'0')
    for tag,val in [('start','1'),('numFmt','decimal'),('lvlText','%1.'),('lvlJc','left')]:
        e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);lvl.append(e)
    pp=OxmlElement('w:pPr');ind=OxmlElement('w:ind');ind.set(qn('w:left'),'360');ind.set(qn('w:hanging'),'360');pp.append(ind);lvl.append(pp)
    rp=OxmlElement('w:rPr');rf=OxmlElement('w:rFonts')
    rf.set(qn('w:ascii'),'Times New Roman');rf.set(qn('w:hAnsi'),'Times New Roman');rf.set(qn('w:eastAsia'),'宋体')
    rp.append(rf);lvl.append(rp)
    aid.append(lvl);numbering.append(aid)
    numid=max([int(x.get(qn('w:numId'))) for x in numbering.findall(qn('w:num'))]+[0])+1
    num=OxmlElement('w:num');num.set(qn('w:numId'),str(numid))
    ai=OxmlElement('w:abstractNumId');ai.set(qn('w:val'),str(abstract_id));num.append(ai);numbering.append(num)
    return numid


def build(md, output):
    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.7)
    sec.left_margin=sec.right_margin=Inches(.82)
    sec.header_distance=sec.footer_distance=Inches(.3)
    normal=doc.styles['Normal'];fonts(normal,11)
    # Remove decorative rules inherited from the bundled blank document.
    for st in doc.styles:
        for border in list(st.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    pf=normal.paragraph_format
    pf.line_spacing=1.18;pf.space_after=Pt(5);pf.first_line_indent=Pt(22)
    pf.widow_control=True
    for name,size in [('Title',20),('Heading 1',15),('Heading 2',13),('Heading 3',11.5)]:
        fonts(doc.styles[name],size,'宋体',True)
        f=doc.styles[name].paragraph_format
        f.first_line_indent=Pt(0);f.space_before=Pt(12);f.space_after=Pt(6)
        f.keep_with_next=True;f.keep_together=True
    fonts(doc.styles['Caption'],9.5)
    doc.styles['Caption'].paragraph_format.first_line_indent=Pt(0)
    doc.styles['Caption'].paragraph_format.space_after=Pt(6)
    doc.styles['Caption'].font.italic=False
    hp=sec.header.paragraphs[0];hp.paragraph_format.first_line_indent=Pt(0)
    hp.add_run('CLE  实验探索周报').font.size=Pt(9)
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;fp.paragraph_format.first_line_indent=Pt(0)
    fp.add_run('第 ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
    fp.add_run(' 页')
    pending=None;bid=0;numbering_id=None;pagebreak_before=False
    lines=md.read_text(encoding='utf-8-sig').splitlines()
    i=0
    while i<len(lines):
        line=lines[i].strip(); i+=1
        if not line: continue
        if line=='<!-- pagebreak -->':
            pagebreak_before=True;numbering_id=None;continue
        anchor=re.fullmatch(r'<a id="([^"]+)"></a>',line)
        if anchor:pending=anchor[1];continue
        im=re.fullmatch(r'!\[([^\]]+)\]\(([^)]+)\)',line)
        if im:
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            figure_width = 4.3 if im[2].endswith('a9_q_learning.png') else 5.0 if im[2].endswith('a10_fs_progress.png') else 6.75
            run=p.add_run();run.add_picture(str(md.parent/im[2]),width=Inches(figure_width))
            dp=run._r.xpath('.//wp:docPr')[0];dp.set('descr',im[1])
            cap=doc.add_paragraph(im[1],style='Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
            continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i].strip());i+=1
            rows=[r for r in rows if not re.fullmatch(r'[| :\-]+',r)]
            cells=[[s.strip() for s in r.strip('|').split('|')] for r in rows]
            tb=doc.add_table(rows=0,cols=len(cells[0]));tb.alignment=WD_TABLE_ALIGNMENT.CENTER
            tb.autofit=False
            borders=OxmlElement('w:tblBorders')
            for side in ('top','left','bottom','right','insideH','insideV'):
                edge=OxmlElement('w:'+side)
                for key,value in (('val','single'),('sz','4'),('color','D9D9D9')):
                    edge.set(qn('w:'+key),value)
                borders.append(edge)
            tb._tbl.tblPr.append(borders)
            n=len(cells[0]);ratios={3:[.26,.30,.44],4:[.43,.19,.19,.19],5:[.31,.17,.17,.18,.17],6:[.25,.17,.145,.145,.145,.145]}.get(n,[1/n]*n)
            if n == 4 and cells[0][-1] == '结果含义':
                ratios = [.30, .19, .15, .36]
            for col,ratio in zip(tb.columns,ratios): col.width=Inches(6.75*ratio)
            for j,row in enumerate(cells):
                tr=tb.add_row();tr._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                if j==0:tr._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for k,(cell,txt) in enumerate(zip(tr.cells,row)):
                    cell.width=Inches(6.75*ratios[k]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0)
                    p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                    p.paragraph_format.line_spacing=1.08
                    p.paragraph_format.keep_with_next=(j<len(cells)-1)
                    inline(p,txt,md.parent)
                    for r in p.runs:r.font.size=Pt(9.5);r.bold=(j==0)
                    tcpr=cell._tc.get_or_add_tcPr()
                    shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'EDEFF2' if j==0 else ('F7F8FA' if j%2==0 else 'FFFFFF'));tcpr.append(shading)
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
            continue
        hm=re.match(r'^(#{1,4}) (.*)$',line)
        nm=re.match(r'^\d+\. (.*)$',line)
        if hm:
            n=len(hm[1]);style='Title' if n==1 else 'Heading '+str(n-1)
            p=doc.add_paragraph(style=style);inline(p,hm[2],md.parent);numbering_id=None
            if n==1:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif nm:
            if numbering_id is None:numbering_id=new_numbering(doc)
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.left_indent=Inches(.25)
            np=OxmlElement('w:numPr')
            lvl=OxmlElement('w:ilvl');lvl.set(qn('w:val'),'0');np.append(lvl)
            ni=OxmlElement('w:numId');ni.set(qn('w:val'),str(numbering_id));np.append(ni)
            p._p.get_or_add_pPr().append(np);inline(p,nm[1],md.parent)
        else:
            p=doc.add_paragraph();inline(p,line,md.parent);numbering_id=None
            if line.startswith(('报告周期：','证据对应：','实验导航：','对应原实验：','数据来源：','[返回','原始证据：')):
                p.paragraph_format.first_line_indent=Pt(0)
            if line.startswith('报告周期：'):
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(9.5)
            if line.startswith(('对应原实验：','数据来源：','[返回')):
                for r in p.runs:r.font.size=Pt(9.5)
            if line == '实验步骤：':
                p.paragraph_format.keep_with_next = True
        if pagebreak_before:
            p.paragraph_format.page_break_before=True
            pagebreak_before=False
        if pending:bid+=1;bookmark(p,pending,bid);pending=None
    doc.core_properties.title=lines[0].removeprefix('# ')
    doc.core_properties.subject='CLE weekly experiment report with requirement-to-evidence navigation'
    doc.core_properties.author='CLE项目研究记录'
    doc.save(output)
    return output


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--template-only',action='store_true')
    parser.add_argument('--source',type=Path)
    parser.add_argument('--output',type=Path)
    args=parser.parse_args()
    if args.source:
        if not args.output:
            parser.error('--source requires --output')
        print(build(args.source,args.output))
        return
    files=[build(TEMPLATE,TEMPLATE.with_suffix('.docx'))]
    if not args.template_only:
        figures()
        files.append(build(REPORT/'CLE-RPT-20261022-20261028-weekly-v04.md', REPORT/'10月22日-10月28日实验探索报告.docx'))
    print(json.dumps([str(f) for f in files],ensure_ascii=False))


if __name__=='__main__':main()
