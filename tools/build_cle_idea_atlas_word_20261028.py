"""将保存好的完整七阶段思路图逐张汇入Word，不刷新实验状态或重画原图。"""
from pathlib import Path
import json
import re
from urllib.parse import unquote
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT/'docs/research/diagrams/ideas_horizontal'
OUT = ROOT/'docs/research/reports/CLE-IDEAS-ATLAS-20261028-v01'
TITLE = 'CLE各条研究思路图汇总'
TEMPLATE = ROOT/'docs/research/templates/CLE-TPL-weekly-idea-map-v02'
SVG = '{http://www.w3.org/2000/svg}'


def fonts(style, size=None):
    style.font.name = 'Times New Roman'
    if size: style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    f = rpr.find(qn('w:rFonts'))
    if f is None:
        f = OxmlElement('w:rFonts');rpr.insert(0, f)
    for k, v in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('cs','Times New Roman'),('eastAsia','宋体')]:
        f.set(qn('w:'+k),v)
    for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        f.attrib.pop(qn('w:'+k),None)


def setup(doc):
    for style in doc.styles:
        if style.type in (1,2): fonts(style)
    for name, size in [('Normal',11),('Title',20),('Heading 1',15),('Hyperlink',11)]:
        if name in doc.styles: fonts(doc.styles[name],size)
    for name in ['Title','Heading 1']:
        doc.styles[name].font.color.rgb=RGBColor(0,0,0)
        for border in list(doc.styles[name].element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for name in ['Normal','Title','Heading 1']:
        pf=doc.styles[name].paragraph_format
        pf.space_after=Pt(5);pf.line_spacing=1.05
    doc.styles['Heading 1'].paragraph_format.keep_with_next=True
    sec=doc.sections[0];sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width=Cm(29.7);sec.page_height=Cm(21)
    sec.left_margin=sec.right_margin=Cm(1.5)
    sec.top_margin=sec.bottom_margin=Cm(1.2)
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)


def bookmark(p, name, number):
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(number));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(number))
    p._p.insert(0,start);p._p.append(end)


def link(p,label,target):
    h=OxmlElement('w:hyperlink')
    if target.startswith('#'): h.set(qn('w:anchor'),target[1:])
    else:
        h.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    f=OxmlElement('w:rFonts')
    for k,v in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','宋体')]:f.set(qn('w:'+k),v)
    pr.append(f)
    c=OxmlElement('w:color');c.set(qn('w:val'),'24577D');pr.append(c)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u)
    r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)


def file_url(path,anchor=''):
    return path.resolve().as_uri()+('#'+anchor if anchor else '')


def diagram_info(svg):
    root=ET.fromstring(svg.read_text(encoding='utf-8'))
    title=next(e.text for e in root.findall(SVG+'text') if e.text)
    refs={}
    for a in root.iter(SVG+'a'):
        href=a.get('href') or a.get('{http://www.w3.org/1999/xlink}href')
        if not href:continue
        path,_,anchor=unquote(href).partition('#')
        if re.fullmatch(r'[A-D]\d{2}E?',anchor,re.I):
            target=(svg.parent/path).resolve()
            if not target.is_file():raise FileNotFoundError(target)
            refs.setdefault(anchor.upper(), file_url(target,anchor))
    return {'title':title,'svg':svg,'png':svg.with_suffix('.png'),'refs':refs}


def add_map(doc,info,index,template=False):
    sec=doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation=WD_ORIENT.LANDSCAPE
    w,h=Image.open(info['png']).size
    sec.page_width=Cm(55);sec.page_height=Cm(52.5*h/w+5.2)
    sec.left_margin=sec.right_margin=Cm(1.25)
    sec.top_margin=sec.bottom_margin=Cm(.8)
    sec.footer_distance=Cm(.25)
    p=doc.add_paragraph(('图册版式示例　' if template else f'I{index:02}　')+info['title'],'Heading 1')
    bookmark(p,f'idea{index:02}',index+1)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(info['png']),width=Cm(52.5))
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0)
    link(p,'返回目录','#index');p.add_run('　')
    link(p,'打开可点击原图',file_url(info['svg']));p.add_run('　实验链接：')
    for label,target in info['refs'].items():
        link(p,label,target);p.add_run('　')


def make_atlas(infos):
    doc=Document();setup(doc)
    p=doc.add_paragraph(TITLE,'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    bookmark(p,'index',1)
    doc.add_paragraph('本册汇总现有九张七阶段完整思路图。每张图保留问题、阶段、实验节点、状态、结果及全部连线。点击目录进入对应图页，图下实验编号可打开原实验说明。')
    doc.add_paragraph('图中状态沿用原图保存时点。这里汇总跨周研究全貌；每周提交时只选入该周对应的图，并保留当周截止状态。')
    for i,info in enumerate(infos,1):
        p=doc.add_paragraph();link(p,f'I{i:02}　{info["title"]}',f'#idea{i:02}')
    doc.add_paragraph('使用横向大页面和高清原图，便于放大查看分支及实验卡片；图下可打开矢量SVG。')
    for i,info in enumerate(infos,1):add_map(doc,info,i)
    doc.core_properties.title=TITLE;doc.core_properties.author='CLE项目研究记录'
    path=OUT/(TITLE+'.docx');doc.save(path)
    lines=['# '+TITLE,'','汇总已有九张完整阶段思路图；Word每张独立一页，目录可跳转。原图中的节点、阶段、状态、结果和连线原样保留。','',f'[下载Word图册]({TITLE}.docx)','']
    for i,info in enumerate(infos,1):
        rel='../../diagrams/ideas_horizontal/'+info['svg'].name
        lines += [f'## I{i:02} {info["title"]}','',f'![{info["title"]}]({rel[:-4]}.png)','',f'[打开可点击原图]({rel})','']
    (OUT/(TITLE+'.md')).write_text('\n'.join(lines),encoding='utf-8')


def make_template(info):
    doc=Document();setup(doc)
    p=doc.add_paragraph('每周思路图汇总模板','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    bookmark(p,'index',1)
    for text in [
        '提交物是一份汇总完整思路图的横版Word图册。截图中的七阶段图是排版依据，每张独立一页；完整保留阶段栏、实验节点、编号、状态、结果、分支与连线。',
        '先列可点击目录，再按研究问题逐张放入原图。每张图下提供返回目录、可点击原图及实验编号链接。Word嵌入完整高清原图，矢量SVG另附链接。',
        '周报选入对应周的思路图，保存该周统计截止时点；跨周全系列另作总图册。历史周不能直接采用包含后续结果的当前全局图。',
        '页面横向，按完整图宽度设置足够大的纸张；保证节点文字可放大清晰阅读。汉字宋体，字母数字Times New Roman。以下一页仅展示完整原图的汇入方式。',
    ]:doc.add_paragraph(text)
    p=doc.add_paragraph();link(p,'查看完整原图汇入示例','#idea01')
    add_map(doc,info,1,template=True)
    doc.core_properties.title='每周思路图汇总模板';doc.save(TEMPLATE.with_suffix('.docx'))
    TEMPLATE.with_suffix('.md').write_text('''# 每周思路图汇总模板

［本周日期范围］思路图汇总

［点击目录进入各条思路图］

## ［图I01 研究问题］

![完整七阶段思路图](［对应周原图.png］)

［返回目录］　［打开可点击SVG原图］　［实验编号及详细说明链接］

每条思路独立一页，汇成一份横版Word。原样保留阶段栏、实验卡片、编号、状态、结果与全部分支连线；按原图汇总，不另改为四栏摘要图或文字表格。每周只使用该周截止时的图，跨周全系列另作总图册。Word嵌入完整高清原图，矢量SVG另附链接；纸张宽度按图调整，汉字宋体、字母数字Times New Roman。示例与生成入口见[完整思路图册](../reports/CLE-IDEAS-ATLAS-20261028-v01/CLE各条研究思路图汇总.md)和 `tools/build_cle_idea_atlas_word_20261028.py`。
''',encoding='utf-8')


if __name__=='__main__':
    OUT.mkdir(parents=True,exist_ok=True)
    infos=[diagram_info(p) for p in sorted(SOURCE.glob('*.svg'))]
    if len(infos)!=9:raise ValueError('应汇总既有九张完整图')
    make_atlas(infos);make_template(infos[-1])
    print(json.dumps({'maps':len(infos),'atlas':str(OUT/(TITLE+'.docx')),'template':str(TEMPLATE.with_suffix('.docx'))},ensure_ascii=False))
