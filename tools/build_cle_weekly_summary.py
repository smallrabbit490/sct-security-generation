"""生成500字以内的周报配套摘要和一页Word；只读已有报告，不调用模型API。"""
import argparse
import csv
import json
import os
from pathlib import Path
from urllib.parse import quote

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
REPORTS=ROOT/'docs/research/reports'
FONT='C:/Windows/Fonts/simsun.ttc'


def font(style,size):
    style.font.name='Times New Roman';style.font.size=Pt(size)
    style.font.color.rgb=RGBColor(0,0,0)
    props=style.element.get_or_add_rPr()
    fonts=props.find(qn('w:rFonts'))
    if fonts is None:
        fonts=OxmlElement('w:rFonts');props.insert(0,fonts)
    for k,v in [('eastAsia','宋体'),('ascii','Times New Roman'),('hAnsi','Times New Roman'),('cs','Times New Roman')]:
        fonts.set(qn('w:'+k),v)
    for name in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        fonts.attrib.pop(qn('w:'+name),None)


def overview(data,path):
    """概览图中的颜色仅区分研究判断，不暗示进度或效应大小。"""
    scale=3
    image=Image.new('RGB',(960*scale,330*scale),'white');draw=ImageDraw.Draw(image)
    def text(x,y,value,size=20,color='#1e293b'):
        draw.multiline_text((x*scale,y*scale),value,font=ImageFont.truetype(FONT,size*scale),fill=color,spacing=7*scale,align='center')
    def box(bounds,fill,outline):
        draw.rounded_rectangle(tuple(int(v*scale) for v in bounds),radius=11*scale,fill=fill,outline=outline,width=scale)
    box((8,7,952,62),'#edf3f8','#c9d8e5')
    text(32,23,data['question'],21)
    fills=[('#edf5ef','#b5d0bc'),('#fff4e5','#e5c79a'),('#edf3fa','#b9cce0')]
    for i,(topic,judgment) in enumerate(data['cards']):
        x=8+i*320
        box((x,94,x+304,227),*fills[i])
        text(x+18,113,topic,20)
        text(x+18,158,judgment,22)
        draw.line((int((x+152)*scale),63*scale,int((x+152)*scale),90*scale),fill='#94a3b8',width=2*scale)
        draw.polygon([(int((x+148)*scale),84*scale),(int((x+156)*scale),84*scale),(int((x+152)*scale),91*scale)],fill='#94a3b8')
    box((8,259,952,320),'#f3f4f6','#d1d5db')
    text(26,278,data['next_figure'],20)
    image.save(path,dpi=(240,240))


def build(data):
    period=data['period'];out=REPORTS/f'CLE-SUM-{period}-v01';out.mkdir(parents=True,exist_ok=True)
    graph_text=[data['question'],data['next_figure']]+[v for pair in data['cards'] for v in pair]
    visible=data['title']+''.join(data['paragraphs'])+''.join(graph_text)+'本周研究主线与主要判断'+'查看本周完整实验报告'
    count=sum(not c.isspace() for c in visible)
    if count>500:raise ValueError(f'summary_exceeds500:{period}:{count}')
    for source in data['evidence']:
        if not (ROOT/source).is_file():raise FileNotFoundError(source)
    overview(data,out/'overview.png')
    source_md=ROOT/data['report_md'];source_word=source_md.with_name(data['report_docx'])
    if not source_word.is_file():raise FileNotFoundError(source_word)
    md_link=Path(os.path.relpath(source_md,out)).as_posix()
    md=out/(data['title']+'.md')
    md.write_text('# '+data['title']+'\n\n'+'\n\n'.join(data['paragraphs'])+
        '\n\n![本周研究主线与主要判断](overview.png)\n\n[查看本周完整实验报告]('+md_link+')\n',encoding='utf-8')
    doc=Document();sec=doc.sections[0]
    # 某些Word默认模板的标题带段落横线；摘要采用无装饰报告标题。
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.1);sec.bottom_margin=Cm(2.0);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
    for name,size in [('Normal',12),('Title',18),('Caption',10)]:font(doc.styles[name],size)
    doc.styles['Caption'].font.bold=False
    doc.styles['Caption'].font.italic=False
    for name in ('Hyperlink','Default Paragraph Font'):
        if name in doc.styles:font(doc.styles[name],11)
    doc.styles['Normal'].paragraph_format.line_spacing=1.3
    doc.styles['Normal'].paragraph_format.space_after=Pt(9)
    p=doc.add_paragraph(data['title'],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(18)
    for value in data['paragraphs']:
        p=doc.add_paragraph(value);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent=Pt(24)
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(10)
    p.add_run().add_picture(str(out/'overview.png'),width=Cm(16.7))
    pic=doc.inline_shapes[-1]._inline.docPr;pic.set('descr','；'.join(graph_text))
    p=doc.add_paragraph('本周研究主线与主要判断','Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(9)
    link=OxmlElement('w:hyperlink')
    target=quote(Path(os.path.relpath(source_word,out)).as_posix(),safe='/.-')
    link.set(qn('r:id'),doc.part.relate_to(target,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');props=OxmlElement('w:rPr');rf=OxmlElement('w:rFonts')
    for k,v in [('eastAsia','宋体'),('ascii','Times New Roman'),('hAnsi','Times New Roman')]:rf.set(qn('w:'+k),v)
    props.append(rf);run.append(props);text=OxmlElement('w:t');text.text='查看本周完整实验报告';run.append(text);link.append(run);p._p.append(link)
    word=out/(data['title']+'.docx');doc.save(word)
    manifest={**data,'visible_characters_including_figure':count,'limit':500,'API_calls':0,
        'output_md':str(md.relative_to(ROOT)),'output_docx':str(word.relative_to(ROOT)),
        'full_reports_modified':False,'word_and_markdown_same_text':True,'chart_kind':'qualitative_research_judgment_not_numeric_effect'}
    (out/'summary.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    return {'period':period,'characters':count,'word':str(word),'markdown':str(md)}


def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('source',type=Path);args=parser.parse_args()
    data=json.loads(args.source.read_text(encoding='utf-8-sig'))
    print(json.dumps([build(item) for item in data],ensure_ascii=False))


if __name__=='__main__':main()
