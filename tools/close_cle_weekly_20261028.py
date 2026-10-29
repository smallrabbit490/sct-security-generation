from pathlib import Path
import json,re,shutil,os,urllib.parse,sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
import build_cle_weekly_report as render
import build_cle_weekly_package_20261028 as pkg
import build_cle_idea_atlas_word_20261028 as atlas
import build_cle_horizontal_ideas as diagrams
R=Path(__file__).resolve().parents[1];O=R/'docs/research/reports/CLE-PKG-20261022-20261028-closeout-v01';OLD=O.with_name('CLE-RPT-20261022-20261028-weekly-v07');S=json.loads((O/'evidence/runtime_snapshot.json').read_text(encoding='utf-8-sig'));P={x['key']:x for x in S['packages']};Q=P['FQ_LEARN'];F=P['FFS'];cut='10月28日23:01'
shutil.copytree(OLD/'figures',O/'figures',dirs_exist_ok=True)
# Preserve original native summaries, read-only; live values may advance after the cutoff.
checks=[]
for x in [Q,F]:
 p=Path(x['segment'])/'summary.json';native=json.loads(p.read_text(encoding='utf-8-sig'));(O/'evidence'/f"native_{x['key']}.json").write_text(json.dumps(native,ensure_ascii=False,indent=2),encoding='utf-8');checks.append({'package':x['key'],'path':str(p),'native_completed':native['completed'],'snapshot_completed':x['summary']['completed'],'native_updated':native.get('updated_utc')})
(O/'evidence/source_checks.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
s=(OLD/(OLD.name+'.md')).read_text(encoding='utf-8')
s=s.replace('10月28日18:10','10月28日23:01').replace('18:10物理账本','23:01物理账本').replace('894/1122','1092/1122').replace('1952/2980','2315/2980').replace('完成45题','完成243题').replace('713/1939','762/1939').replace('713张','762张')
s=s.replace('全量主比较保留，中间探索按信息价值减量。','已启动全量继续，其他思路先做探针和分析，排队全量暂缓。')
s=s.replace('再对298题做四组两重复','先做2题四组两重复的技术探针；298题四组两重复暂缓').replace('再决定追加一批或进入全量策略比较','再决定追加一批或进入6题策略探针').replace('保留已启动全量和最终主要比较，缩小中间回测、标签和检索批次','保留已启动全量，其余先接探针；排队全量先跳过').replace('当前读数不构成分栏收益，最终按全部同题配对结果判定。','当前各条件已测集合不同，最终按全部同题配对结果判定。')
s=s.replace('原卡阶段读数173/390，分栏169/390','原卡阶段读数198/462，分栏194/464').replace('原混合卡173/390，分栏169/390','原混合卡198/462，分栏194/464')
# Replace todo steps, retaining the scientific criteria and origins in the weekly source.
s=s.replace('工作台按前置结果接续，不因普通负结果取消已保留的主比较。','工作台按真实前置接续探针和分析；排队全量暂缓，不能因探针通过自动启动。')
s=s.replace('41个执行节点已经安装，61项离线验证通过。','探针接力规则v09已安装；40项离线调度验证通过。')
s=s.replace('新库完成后，在Base115题和Plus140题分别做无卡、旧库、新库两重复，共1530次。','新库与原环境就绪后先登记各集2题的有界检查；1530次全量保持暂缓。')
s=s.replace('K3和相似度另覆盖全部273题，合计1416次。','五条件合计540次；原1416次中的其余876次暂缓。')
s=s.replace('完成分析后，298题做原CLE、相似度、学习Q、同CWE打乱Q四组两重复，共2384次；先固定异反馈题族子集。','完成前置分析后先2题四组两重复，共16次技术探针；2384次计划中的其余部分暂缓，异反馈题族子集事先固定。')
s=s.replace('全273题比较回答最终评分价值。','全273题最终比较仍待后续启动决定。')
s=s.replace('再进行298题三策略两重复，共1788个策略结果。','再先做6题三策略两重复，共36次；1788个策略结果的其余部分暂缓。')
s=s.replace('从空库按原CLE初始化、三轮主动回放、改卡和全局门控建立新库；不是机械重跑三遍273题。','从空库先做2道初始化题验证建库流程；完整初始化、三轮主动回放和全局门控暂缓，两题产物不冒充完整库。')
s=s.replace('在298题上比较v4.1无卡、旧库、新库，两重复共1788个结果；1192个无卡和旧库结果逐格核准复用，新生成最多596次。','真实完整新库到位后，先12题三条件两重复，共72次；48次精确核准复用，新生成最多24次，全298题暂缓。')
# Rebuild the two live numerical tables from one cutoff.
start=s.index('| 学习阶段 |');end=s.index('\n\n',start)
t=['| 学习阶段 | 已完成/计划 | 可测 | 首次Joint | 最终Joint | 已获反馈卡/1939 |','|---|---:|---:|---:|---:|---:|']
for k,x in Q['stages'].items():
 a=x['final'];t.append(f"| {k} | {x['completed']}/{x['planned']} | {a['measured']} | {x['first']['joint']} | {a['joint']}（{100*a['joint']/a['measured']:.2f}%） | {x['Q_covered_cards']}（{100*x['Q_coverage']:.2f}%） |")
s=s[:start]+'\n'.join(t)+s[end:]
start=s.index('| 条件 | 已完成 |');end=s.index('\n\n',start)
t=['| 条件 | 已完成 | 可测 | F通过 | S通过 | Joint |','|---|---:|---:|---:|---:|---:|'];names={'B0':'无卡','BM':'原混合卡','BF':'仅功能','BS':'仅安全','BFS':'功能安全分栏'}
for k,x in F['summary']['arms'].items():t.append(f"| {k} {names[k]} | {x['completed']} | {x['measured']} | {x['functional']} | {x['security']} | {x['joint']}（{100*x['joint']/x['measured']:.2f}%） |")
s=s[:start]+'\n'.join(t)+s[end:]
# Keep historical results and actual case analyses; replace the superseded scheduling chapter.
s=s[:s.index('<a id="a11">')]+'''<a id="a11"></a>
### 【A11】怎样用探针接力控制成本？

**结论：已启动Q学习和功能安全全量继续；其他研究先做探针，排队全量先跳过。** v09取代白天v08的自动接续方式，40项调度离线验证通过。这是执行规则验证，不是科学收益结果。

#### 为什么这样安排？

旧Q、分栏和分类器小实验均未胜过对照。当前先用有限样本确认新机制和训练机会，把后续全量投入留给最有价值的问题；已经产生大量数据的两项全量保持规则完成。

#### 怎样执行？

1. 完成活动学习和功能安全，回放保存结果并阅读实际代码。
2. 有真实前置且入口已验证的探针自动接续，普通负结果按既定固定样本收口；条件追加只执行登记的小批。
3. 遇到排队全量先跳过，继续其他独立探针。跳过不解除真实完整卡库、分类器或原镜像依赖。

| 研究任务 | 本次自动范围 | 全量状态 |
|---|---|---|
| Q同题回测 | 54题×5条件×2＝540次 | 其余876次暂缓 |
| Q审计 | 2题×4条件×2＝16次技术探针 | 其余2368次暂缓 |
| 记忆标签 | 82题×5档×2＝820次；条件满足才追加82题 | 不直接扩到273题 |
| 自动选档 | 真实新标签训练后6题×3策略×2＝36次 | 其余1752次暂缓 |
| 检索 | 57题；条件满足才追加57题 | 不直接全571题 |
| 新库 | 2道初始化题技术探针 | 完整建库暂缓 |
| 卡库交叉 | 真实完整新库到位后12题×3×2＝72次 | 全298题暂缓 |
| 局部原因复验 | 最多6题、12次本地验证 | 无实验模型请求 |

新增有界入口仍需离线验证后才能实际派发；队列名称不等于已做实验。截止时这些后续探针没有真实完成行。

#### 已经投入多少？

'''
t=['| 本轮包或阶段 | 已报告Token | 物理HTTP | 未知用量或在途HTTP |','|---|---:|---:|---:|']
for k,x in Q['stages'].items():t.append(f"| Q：{k} | {x['known_reported_tokens']:,} | {x['physical_HTTP']:,} | {x['unknown_or_inflight_HTTP']:,} |")
x=F['summary'];t.append(f"| 功能安全 | {x['known_reported_tokens']:,} | {x['physical_HTTP']:,} | {x['unknown_or_inflight_HTTP']:,} |")
x=S['usage'];t.append(f"| 两包合计 | {x['known_reported_tokens']:,} | {x['physical_HTTP']:,} | {x['unknown_or_inflight_HTTP']:,} |")
s+='\n'.join(t)+'''\n\n本表为23:01累计物理账本，未知用量不计为零；基础实验65,186,659 Token在【A8】单列，Codex分析用量另记。没有当前账单就不折算实付人民币。

#### 对主线有什么影响？

覆盖增加证明反馈触达更多卡，不证明代码更正确。下一步54题配对回测先区分学习轮次的影响；功能安全完整对照回答经验拆分能否兼顾两个目标。82题标签和57题检索负责判定是否存在可训练规律、遗漏发生在哪里，不替代完整迁移结论。

[返回本周主线](#main)　[返回待做安排](#todo)　[当前执行规则](../../CLE-TODO-20261022-20261028-v09.md)　[本次保存快照](evidence/runtime_snapshot.json)
'''
s=s.replace('本报告采用10月28日23:01保存的结果。','本报告为周三收尾版，采用10月28日23:01保存的结果。之后新增实质结果作为本周期增补。')
s=s.replace('| K3部分 | 23/33 | 22/33 | 110 | 129 |','| K3部分 | 124/179 | 125/182 | 662 | 835 |').replace('及冻结全量比较','及54题冻结探针；排队全量暂缓').replace('全273题最终Q与相似度比较负责检验最终价值。','全273题最终Q与相似度比较用于后续总体检验，目前暂缓。').replace('下一步应在同一A7流程、相同来源资格和固定全量名单上继续比较','后续比较须固定A7流程、来源资格与名单；当前先按v09做探针，排队全量暂缓')
name='10月22日-10月28日实验探索报告';md=O/(name+'.md');md.write_text(s,encoding='utf-8')
pkg.bars(O/'figures/a9_q_learning.png','各轮获得反馈的卡片数',['K0','K1','K2','K3部分'],[445,610,702,762],maxval=900,foot='卡库1939张；K3完成243/273题，不是完整轮次成绩。')
a=F['summary']['arms'];pkg.bars(O/'figures/a10_fs_progress.png','各条件阶段Joint通过次数',[names[k] for k in a],[x['joint'] for x in a.values()],maxval=240,foot='各臂已测集合不同；最终按同题配对比较，不以当前总数差归因。')
pkg.bars(O/'figures/summary_key.png','旧Q代码对照：Joint通过次数',['相似度','学习Q','打乱Q'],[13,11,13],maxval=16,foot='各计划48次；打乱Q有1次缺测，不计为已测失败。')
short=f'''# 10月22日-10月28日本周结论与下一步

本周检验经验选择、表达和来源能否改善功能安全共同通过率。统计截至{cut}。

**主要发现：** Q已改变选卡，旧Q代码对照学习Q为11/48，相似度13/48，未提高正确率。分栏9/24低于原卡13/24；分类器21/48低于固定提醒22/48。[对应实验](10月22日-10月28日实验探索报告.docx#a2)

![旧Q对照](figures/summary_key.png)

**主线影响：** 新Q学习1092/1122，反馈覆盖762/1939；功能安全2315/2980。覆盖扩大已经发生，收益仍由同题对照检验。[在线学习](10月22日-10月28日实验探索报告.docx#a9)

**下一步：** 完成两项活动全量并分析；随后54题Q回测、82题标签、57题检索按前置自动接力，排队全量暂缓。[待做安排](10月22日-10月28日实验探索报告.docx#todo)

[完整周报](10月22日-10月28日实验探索报告.docx)　[九条思路图](10月22日-10月28日研究思路图汇总.docx)
'''
sp=O/'10月22日-10月28日本周结论与下一步.md';sp.write_text(short,encoding='utf-8')
original=render.add_hyperlink
def link(p,label,target,base):
 if not target.startswith('#') and not urllib.parse.urlsplit(target).scheme:
  path,_,anchor=target.partition('#');target=(base/path).resolve().as_uri()+('#'+anchor if anchor else '')
 original(p,label,target,base)
render.add_hyperlink=link
for path in [md,sp]:
 render.build(path,path.with_suffix('.docx'));d=Document(path.with_suffix('.docx'));nums=iter(re.findall(r'^(\d+)\. ',path.read_text(encoding='utf-8'),re.M))
 for p in d.paragraphs:
  pp=p._p.pPr
  if pp is not None and pp.numPr is not None:
   pp.remove(pp.numPr);run=p.add_run(next(nums)+'）');p._p.remove(run._r);p._p.insert(1,run._r)
 for table in d.tables:
  for row in table.rows:
   for c in row.cells:
    for p in c.paragraphs:p.paragraph_format.keep_with_next=False
 if path==sp:
  for shape in d.inline_shapes:
   shape.height=int(shape.height*0.65);shape.width=int(shape.width*0.65)
 d.save(path.with_suffix('.docx'))
# Render complete seven-stage diagrams in the report-only directory; no live diagram/state writes.
S['tasks']=S.get('autonomy',{}).get('tasks',[]);snap=O/'evidence/diagram_snapshot.json';snap.write_text(json.dumps(S,ensure_ascii=False),encoding='utf-8')
canonical=diagrams.OUT;diagrams.OUT=O/'figures/ideas';diagrams.TEMP=O/'evidence/diagram_render';export=diagrams.export
orig_simplify=diagrams.simplify
def simplify(specs,items,snapshot):
 items=orig_simplify(specs,items,snapshot)
 edits={'C02':('先做54题五条件比较','探针自动接力；其余全量暂缓','54题×5×2＝540次'),'C03':('先检查四臂选择与生成','2题技术探针；全量暂缓','2题×4×2＝16次'),'C06':('先检查新分类器策略','真实训练完成后做小对照','6题×3×2＝36次'),'C09':('先检查空库建立流程','两题产物不能代替完整库','2道初始化题；全量暂缓'),'C10':('比较新旧库的小样本效果','等待真实完整新库；全量暂缓','12题×3×2＝72次'),'C11':('原环境到位后先小检查','全量暂缓；原镜像仍缺','Base/Plus各前2题')}
 for k,(title,meaning,scale) in edits.items():items[k].update(title=title,meaning=meaning,scale=scale)
 for spec in specs:
  for n in diagrams.walk(spec['root']):
   for identity in n.experiments:
    if identity in edits:n.title,note,scale=edits[identity];n.text=[note,scale]
   if 'C02' in n.experiments:n.title='先做54题回测';n.text=['五条件两重复共540次。','其余876次全量先跳过。']
 return items
diagrams.simplify=simplify
def exp(spec,items):
 export(spec,items);p=diagrams.OUT/(spec['stem']+'.svg');text=p.read_text(encoding='utf-8')
 def rebase(m):
  url=urllib.parse.urlsplit(urllib.parse.unquote(m[1]))
  if url.scheme or not url.path:return m[0]
  target=(canonical/url.path).resolve();return 'href="'+os.path.relpath(target,p.parent).replace('\\','/')+('#'+url.fragment if url.fragment else '')+'"'
 p.write_text(re.sub(r'href="([^"]+)"',rebase,text),encoding='utf-8')
diagrams.export=exp;diagrams.main(snap)
infos=[atlas.diagram_info(p) for p in sorted(diagrams.OUT.glob('*.svg'))];d=Document();atlas.setup(d);title='10月22日-10月28日研究思路图汇总';p=d.add_paragraph(title,'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;atlas.bookmark(p,'index',1);d.add_paragraph('周三收尾，统计截至10月28日23:01。完整七阶段图保留历史证据；活动全量继续，其余先做探针，排队全量暂缓。')
lines=['# '+title,'','统计截至10月28日23:01。与完整周报及一页结论同源。','[完整周报](10月22日-10月28日实验探索报告.docx)','']
for i,info in enumerate(infos,1):
 p=d.add_paragraph();atlas.link(p,f'I{i:02} {info["title"]}',f'#idea{i:02}');lines.extend([f'<a id="idea{i:02}"></a>',f'## I{i:02} {info["title"]}',f'![完整七阶段图](figures/ideas/{info["png"].name})',f'[SVG原图](figures/ideas/{info["svg"].name})',''])
for i,info in enumerate(infos,1):atlas.add_map(d,info,i)
d.save(O/(title+'.docx'));(O/(title+'.md')).write_text('\n'.join(lines),encoding='utf-8')
# Count visible summary text plus visible figure text; URLs and Markdown syntax excluded.
visible=re.sub(r'!?\[([^\]]*)\]\([^)]+\)',r'\1',short);visible=re.sub(r'[#*\s]','',visible);chart='旧Q代码对照：Joint通过次数相似度13学习Q11打乱Q13各计划48次；打乱Q有1次缺测，不计为已测失败。';count=len(visible)+len(chart);assert count<=500,count
(O/'build_manifest.json').write_text(json.dumps({'status':'Wednesday_closeout','cutoff':S['time'],'summary_visible_chars':count,'documents':3,'model_requests':0,'runtime_modified':False,'plan':'v09'},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'summary_chars':count,'output':str(O)},ensure_ascii=False))
